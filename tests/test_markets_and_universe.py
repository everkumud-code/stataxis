import io
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import Session

from api.export import export_live_stats_xlsx
from api.media_intelligence import market_report
from collector.run import run_collection_pass
from collector.storage import Base, Channel, CollectionRun, Observation, Video, create_database, save_observations
from collector.youtube.client import YouTubeAPIError
from collector.youtube.collector import ChannelTarget
from metrics.markets import market_label, normalize_segment
from tests.test_collection_run import FakeClient

AS_OF = datetime(2026, 9, 16, 10, 0, tzinfo=UTC)


def test_market_labels():
    assert market_label("news", "Hindi") == "Hindi News"
    assert market_label(None, "Tamil") == "Tamil News"
    assert market_label("business", "English") == "Business News"
    assert market_label("business", "Hindi") == "Business News"
    assert market_label("print", "Hindi") == "Print Media"
    assert market_label("party", "Hindi") == market_label("leader", "English") == "Political Parties & Leaders"
    assert market_label("commentator", "Hindi") == "Political Commentators"
    assert market_label("news", "unknown") == "News"
    assert normalize_segment("") == "news" and normalize_segment(" Business ") == "business"
    with pytest.raises(ValueError):
        normalize_segment("sports")


def _seed_report():
    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    session = Session(engine)
    spec = [("a", "Alpha", "Hindi", "news", 1000, 1600), ("b", "Bravo", "Hindi", "news", 2000, 2200),
            ("c", "Charlie", "English", "business", 100, 300), ("d", "Delta", "Hindi", "business", 500, 700),
            ("e", "Echo", "Hindi", "print", 10, 10)]
    for key, name, language, segment, start, end in spec:
        channel = Channel(youtube_channel_id=key, name=name, language=language, region="India", segment=segment)
        session.add(channel)
        session.flush()
        video = Video(youtube_video_id=key + "v", channel_id=channel.id, title=name, published_at="2026-09-10T09:00:00+00:00")
        session.add(video)
        session.flush()
        for offset, views in ((timedelta(minutes=50), start), (timedelta(0), end)):
            session.add(Observation(video_id=video.id, channel_id=channel.id, observed_at=AS_OF - offset, view_count=views,
                                    is_live=False, classification="REGULAR_VIDEO"))
    session.commit()
    return session


def test_market_report_share_percent_headline_and_per_market_summary():
    payload = market_report(_seed_report(), as_of=AS_OF, period="1h")
    rows = {row["channel"]: row for row in payload["channels"]}
    assert rows["Alpha"]["share_percent"] == 75.0 and rows["Bravo"]["share_percent"] == 25.0
    assert rows["Charlie"]["share_percent"] == 50.0 and rows["Delta"]["share_percent"] == 50.0
    assert rows["Alpha"]["market_label"] == "Hindi News" and rows["Charlie"]["market_label"] == "Business News"
    assert rows["Alpha"]["share_basis"] == "view_delta"
    markets = {market["label"]: market for market in payload["markets"]}
    assert set(markets) == {"Hindi News", "Business News", "Print Media"}
    assert markets["Hindi News"]["headline"] == "Hindi News viewership (1h): Alpha led with 75% share followed by Bravo"
    assert markets["Business News"]["leaders"][0]["share_percent"] == 50.0
    # Print has a single channel whose gain is zero: no share can be computed, and nothing is invented.
    assert rows["Echo"]["share_percent"] is None
    assert markets["Print Media"]["headline"] is None


def test_market_report_segment_and_market_filters():
    session = _seed_report()
    business = market_report(session, as_of=AS_OF, period="1h", segment="business")
    assert {row["channel"] for row in business["channels"]} == {"Charlie", "Delta"}
    assert business["filters"]["segment"] == "business"
    by_label = market_report(session, as_of=AS_OF, period="1h", market="Hindi News")
    assert {row["channel"] for row in by_label["channels"]} == {"Alpha", "Bravo"}
    legacy = market_report(session, as_of=AS_OF, period="1h", market="Hindi / India")  # old "Language / Region" name still works
    assert {row["channel"] for row in legacy["channels"]} >= {"Alpha", "Bravo"}
    with pytest.raises(ValueError, match="segment must be one of"):
        market_report(session, as_of=AS_OF, period="1h", segment="sports")


def test_segment_is_persisted_and_existing_databases_are_migrated(tmp_path):
    engine = create_database("sqlite://")
    with Session(engine) as session:
        save_observations(session, "Zee Business", "UCbiz", "Zee", "Hindi", [], segment="business")
        save_observations(session, "Zee Business", "UCbiz", "Zee", "Hindi", [])  # later call without segment must not reset it
        assert session.query(Channel).one().segment == "business"
        save_observations(session, "Plain", "UCplain", "n", "Hindi", [])
        assert session.query(Channel).filter_by(youtube_channel_id="UCplain").one().segment == "news"
    url = f"sqlite:///{tmp_path / 'old.db'}"
    old = create_engine(url)
    with old.begin() as connection:
        connection.execute(text("CREATE TABLE stx_users (id INTEGER PRIMARY KEY, email VARCHAR(320))"))
        connection.execute(text("CREATE TABLE stx_channels (id INTEGER PRIMARY KEY, youtube_channel_id VARCHAR(64), name VARCHAR(255))"))
        connection.execute(text("INSERT INTO stx_channels (youtube_channel_id, name) VALUES ('UCold', 'Old')"))
        connection.execute(text("CREATE TABLE stx_videos (id INTEGER PRIMARY KEY, youtube_video_id VARCHAR(32), channel_id INTEGER, title TEXT, published_at VARCHAR(64))"))
    migrated = create_database(url)
    with migrated.connect() as connection:
        assert connection.execute(text("SELECT segment FROM stx_channels WHERE youtube_channel_id='UCold'")).scalar() == "news"


def test_channel_target_carries_segment_from_config():
    target = ChannelTarget(**{"channel_id": "UC1", "name": "X", "language": "Hindi", "segment": "print"})
    assert target.segment == "print"
    assert ChannelTarget("UC2", "Y").segment == "news"


class PartlyBrokenClient(FakeClient):
    """UC-gone is not returned by YouTube, UC-boom fails with a plain error; everything else works."""

    def get_channels(self, channel_ids: list[str]):
        return [item for item in super().get_channels(channel_ids) if item["id"] != "UC-gone"]

    def list_uploads(self, uploads_id: str, max_results: int):
        if uploads_id == "uploads-UC-boom":
            raise ValueError("bad payload")
        return super().list_uploads(uploads_id, max_results)


def test_one_bad_channel_does_not_stop_the_rest_of_the_universe():
    engine = create_database("sqlite://")
    with Session(engine) as session:
        targets = [ChannelTarget("UC-gone", "Gone"), ChannelTarget("UC-ok1", "One", segment="business"),
                   ChannelTarget("UC-boom", "Boom"), ChannelTarget("UC-ok2", "Two")]
        result = run_collection_pass(session, PartlyBrokenClient(), targets, max_videos=1)
        assert result.channel_errors == 2
        assert result.videos_observed == 2
        run = session.get(CollectionRun, result.run_id)
        assert run.status == "partial" and "Gone" in run.error_message and "Boom" in run.error_message
        assert {channel.name for channel in session.query(Channel)} == {"One", "Two"}
        assert session.query(Channel).filter_by(name="One").one().segment == "business"


def test_quota_errors_still_abort_the_pass_so_the_service_can_back_off():
    class QuotaClient(FakeClient):
        def get_channels(self, channel_ids: list[str]):
            raise YouTubeAPIError("quota exceeded", status_code=403)

    engine = create_database("sqlite://")
    with Session(engine) as session:
        with pytest.raises(YouTubeAPIError):
            run_collection_pass(session, QuotaClient(), [ChannelTarget("UC-a", "A"), ChannelTarget("UC-b", "B")], max_videos=1)
        assert session.query(CollectionRun).one().status == "failed"


def test_live_stats_excel_export_has_feed_columns_and_method_sheet():
    from tests.test_live_stats import T0, _channel, _sample, _session, _stream
    session = _session()
    channel = _channel(session, "c1", "Alpha")
    main = _stream(session, channel, "main", T0 - timedelta(days=45))
    for minute in range(0, 3):
        _sample(session, main, T0 + timedelta(minutes=minute), 1500)
    session.commit()
    body = export_live_stats_xlsx(session, start_at=T0, end_at=T0 + timedelta(minutes=2))
    workbook = load_workbook(io.BytesIO(body))
    assert workbook.sheetnames == ["Channels", "Markets", "Method"]
    header = [cell.value for cell in workbook["Channels"][1]]
    assert {"All Avg", "All Peak", "Primary Avg", "Secondary Peak", "Share % (All Avg)"} <= set(header)
    assert workbook["Channels"][2][1].value == "Hindi News"
    assert workbook["Channels"][2][4].value == 1500
