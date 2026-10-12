from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from api.export import MAX_EXPORT_ROWS, ObservationExportFilters, export_observations_xlsx, export_response, filtered_observations
from collector.storage import Base, Channel, Observation, Video
from metrics.persistence import IntelligenceSnapshotRecord


def make_session() -> Session:
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return Session(engine)


def test_filtered_observations_respects_filters() -> None:
    session = make_session()
    channel = Channel(youtube_channel_id="UC123", name="Test", network="N", language="Hindi", region="North")
    session.add(channel)
    session.flush()
    video = Video(youtube_video_id="dQw4w9WgXcQ", channel_id=channel.id, title="Test")
    session.add(video)
    session.flush()
    session.add(Observation(video_id=video.id, channel_id=channel.id, observed_at=datetime(2026, 9, 14, 10, tzinfo=UTC), view_count=100, is_live=True, classification="LIVE"))
    session.add(Observation(video_id=video.id, channel_id=channel.id, observed_at=datetime(2026, 9, 14, 12, tzinfo=UTC), view_count=200, is_live=False, classification="VOD"))
    session.commit()

    rows = filtered_observations(session, ObservationExportFilters(
        start_at=datetime(2026, 9, 14, 9, tzinfo=UTC),
        end_at=datetime(2026, 9, 14, 11, tzinfo=UTC),
        language="Hindi", region="North", is_live=True,
    ))
    assert len(rows) == 1
    assert rows[0][0].view_count == 100


def test_export_contains_two_sheets_and_latest_intelligence() -> None:
    session = make_session()
    channel = Channel(youtube_channel_id="UC456", name="Test", network="N", language="English", region="South")
    session.add(channel)
    session.flush()
    video = Video(youtube_video_id="aBcDeFgHiJk", channel_id=channel.id, title="Headline")
    session.add(video)
    session.flush()
    observed = datetime(2026, 9, 14, 10, tzinfo=UTC)
    session.add(Observation(video_id=video.id, channel_id=channel.id, observed_at=observed, view_count=500, is_live=False, classification="VOD"))
    session.add(IntelligenceSnapshotRecord(video_id=video.id, generated_at=observed, score=72.5, confidence=100, available_signals=7, view_json='{"view":"Positive evidence"}', contributions_json="[]"))
    session.commit()

    payload = export_observations_xlsx(session, ObservationExportFilters(language="English"))
    workbook = load_workbook(BytesIO(payload), read_only=True)
    assert workbook.sheetnames == ["StatAxis Data", "StatAxis Intelligence", "Notes"]
    assert workbook["StatAxis Data"].max_row == 2
    assert workbook["StatAxis Intelligence"].cell(2, 4).value == 72.5
    assert workbook["StatAxis Intelligence"].cell(2, 8).value == "Positive evidence"
    # every time in the workbook is Indian time and says so: 10:00 UTC is 3:30 pm in India
    data = workbook["StatAxis Data"]
    assert data.cell(1, 1).value == "Observed At (IST)" and data.cell(2, 1).value == datetime(2026, 9, 14, 15, 30)
    assert workbook["StatAxis Intelligence"].cell(1, 7).value == "Generated At (IST)" and workbook["StatAxis Intelligence"].cell(2, 7).value == datetime(2026, 9, 14, 15, 30)


def test_filtered_observations_rejects_more_than_export_cap() -> None:
    class Result:
        def all(self):
            return [object()] * (MAX_EXPORT_ROWS + 1)

    class FakeSession:
        def execute(self, statement):
            assert statement._limit_clause.value == MAX_EXPORT_ROWS + 1
            return Result()

    try:
        filtered_observations(FakeSession(), ObservationExportFilters())
    except ValueError as exc:
        assert str(exc) == "report has more than 50000 rows; narrow the date range or add filters"
    else:
        raise AssertionError("oversized export was not rejected")


def test_export_response_returns_400_for_oversized_report(monkeypatch) -> None:
    message = "report has more than 50000 rows; narrow the date range or add filters"
    monkeypatch.setattr("api.export.export_for_role", lambda *args, **kwargs: (_ for _ in ()).throw(ValueError(message)))
    status, headers, body = export_response(object(), "PRO", ObservationExportFilters())
    assert status == 400
    assert headers == {"Content-Type": "application/json"}
    assert body.decode() == '{"error": "report has more than 50000 rows; narrow the date range or add filters"}'


def test_ist_cell_turns_any_kind_of_time_into_indian_time_without_a_zone() -> None:
    from api.stamps import ist_cell
    assert ist_cell(datetime(2026, 10, 12, 11, 45, tzinfo=UTC)) == datetime(2026, 10, 12, 17, 15)
    assert ist_cell(datetime(2026, 10, 12, 11, 45)) == datetime(2026, 10, 12, 17, 15)                  # the database keeps naive times as UTC
    assert ist_cell("2026-10-12T11:45:00+00:00") == datetime(2026, 10, 12, 17, 15) and ist_cell("2026-10-12T11:45:00Z") == datetime(2026, 10, 12, 17, 15)
    assert ist_cell(None) is None and ist_cell("") is None and ist_cell("not a time") == "not a time" and ist_cell(7) == 7
