import io
import json
from datetime import timedelta
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

import web
from api.errors import AuthorizationError
from api.live_http import live_application
from collector.storage import Base
from tests.test_live_stats import T0, _channel, _sample, _stream


@pytest.fixture()
def engine():
    engine = create_engine("sqlite://", future=True)
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        channel = _channel(session, "c1", "Alpha")
        main = _stream(session, channel, "main", T0 - timedelta(days=45))
        event = _stream(session, channel, "event", T0 - timedelta(hours=1))
        for minute in range(0, 4):
            _sample(session, main, T0 + timedelta(minutes=minute), 1000)
        _sample(session, event, T0 + timedelta(minutes=2), 250)
        session.commit()
    return engine


def _call(app, path, query=""):
    captured = {}
    environ = {"REQUEST_METHOD": "GET", "PATH_INFO": path, "QUERY_STRING": query, "HTTP_AUTHORIZATION": "Bearer test"}
    body = app(environ, lambda status, headers: captured.update(status=status, headers=dict(headers)))
    return captured["status"], captured["headers"], body[0]


@pytest.fixture()
def app(engine, monkeypatch):
    monkeypatch.setattr("api.live_http.authenticate", lambda header: SimpleNamespace(user_id=1))
    monkeypatch.setattr("api.live_http.policy_for_identity", lambda identity: SimpleNamespace())
    monkeypatch.setattr("api.live_http.require_capability", lambda policy, name: None)
    return live_application(lambda: Session(engine))


def test_web_routes_the_new_live_stats_paths(monkeypatch):
    routed = []

    def fake_live_api(environ, start_response):
        routed.append(environ["PATH_INFO"])
        start_response("200 OK", [("Content-Type", "text/plain")])
        return [b"ok"]

    monkeypatch.setattr(web, "_live_api", fake_live_api)
    paths = ["/api/v1/audience/live/stats", "/api/v1/audience/live/stats/export", "/api/v1/audience/live/snapshot"]
    for path in paths:
        status, _, _ = _call(web.application, path)
        assert status == "200 OK"
    assert routed == paths


def test_stats_endpoint_returns_feed_split_over_the_window(app):
    query = f"start={T0.isoformat().replace('+', '%2B')}&end={(T0 + timedelta(minutes=3)).isoformat().replace('+', '%2B')}"
    status, headers, body = _call(app, "/api/v1/audience/live/stats", query)
    assert status == "200 OK"
    payload = json.loads(body)
    row = payload["channels"][0]
    assert row["feeds"]["primary"]["average"] == 1000.0
    assert row["feeds"]["secondary"]["peak"] == 250
    assert row["feeds"]["all"]["peak"] == 1250


def test_snapshot_endpoint(app):
    at = (T0 + timedelta(minutes=2)).isoformat().replace("+", "%2B")
    status, _, body = _call(app, "/api/v1/audience/live/snapshot", f"at={at}")
    assert status == "200 OK"
    assert json.loads(body)["channels"][0]["feeds"] == {"primary": 1000, "secondary": 250, "all": 1250}


def test_export_endpoint_returns_an_excel_file(app):
    query = f"start={T0.isoformat().replace('+', '%2B')}&end={(T0 + timedelta(minutes=3)).isoformat().replace('+', '%2B')}"
    status, headers, body = _call(app, "/api/v1/audience/live/stats/export", query)
    assert status == "200 OK"
    assert "spreadsheetml" in headers["Content-Type"] and "attachment" in headers["Content-Disposition"]
    assert load_workbook(io.BytesIO(body)).sheetnames == ["Channels", "Markets", "Method"]


@pytest.mark.parametrize("query", ["bucket_seconds=5", "bucket_seconds=abc", "segment=bogus", "start=not-a-date", "start=2026-09-20T10:00:00%2B00:00&end=2026-09-19T10:00:00%2B00:00"])
def test_bad_parameters_are_400_not_500(app, query):
    status, _, body = _call(app, "/api/v1/audience/live/stats", query)
    assert status == "400 Bad Request"
    assert "error" in json.loads(body)


def test_plan_without_live_access_gets_403(engine, monkeypatch):
    monkeypatch.setattr("api.live_http.authenticate", lambda header: SimpleNamespace(user_id=1))
    monkeypatch.setattr("api.live_http.policy_for_identity", lambda identity: SimpleNamespace())

    def deny(policy, name):
        raise AuthorizationError("upgrade required")

    monkeypatch.setattr("api.live_http.require_capability", deny)
    app = live_application(lambda: Session(engine))
    for path in ("/api/v1/audience/live/stats", "/api/v1/audience/live/snapshot", "/api/v1/audience/live/stats/export"):
        status, _, _ = _call(app, path)
        assert status == "403 Forbidden", path
