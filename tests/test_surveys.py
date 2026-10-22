"""Polls and surveys: the definition and its locks, checking answers, storing them without personal data, duplicates and bots, the results (and when they may be public), exports, invitations, the routes and the pages.
Invented data in a scratch database."""
import io
import json
import re
import shutil
import subprocess
import tempfile
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from api import survey_core as core
from api import survey_pages
from api.auth import AuthIdentity, hash_password, issue_token
from api.rate_limit import limiter
from api.survey_http import survey_admin_application, survey_public_application, survey_site_response
from collector.storage import Base, Survey, SurveyResponse, User

NOW = datetime(2026, 10, 22, 12, 0, tzinfo=UTC)


@pytest.fixture(autouse=True)
def isolate(monkeypatch):
    monkeypatch.setenv("STAXIS_AUTH_SECRET", "test-secret")
    limiter.reset()


def make():
    engine = create_engine("sqlite://", future=True, poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    return Session(engine), engine


def opts(*texts):
    return [{"text": t} for t in texts]


def base(**over):
    q = [{"type": "single", "text": "Do you agree?", "required": True, "options": opts("Yes", "No", "Can't say")},
         {"type": "multi", "text": "Why? (up to 2)", "max_select": 2, "options": opts("Price", "Quality", "Habit", "None of these") + [{"text": "Other", "other": True}], "show_if": None},
         {"type": "rating", "text": "How happy?", "scale_max": 5},
         {"type": "nps", "text": "Recommend?"},
         {"type": "text", "text": "Anything else?"}]
    return {"title": "Test survey", "status": "live", "questions": q, **over}


def live(s, **over):
    return core.save_survey(s, "owner", base(**over))


def answer(s, row, answers, **kw):
    return core.submit_response(s, row, answers, token=kw.pop("token", secrets_token()), now=kw.pop("now", NOW), started_ms=kw.pop("started_ms", (NOW - timedelta(seconds=90)).timestamp() * 1000), **kw)


def secrets_token():
    import secrets
    return secrets.token_hex(12)


# ------------------------------------------------------------------------------------------------------------------------------ the definition
def test_questions_get_stable_ids_and_the_starters_are_all_valid():
    s, _ = make()
    saved = live(s)
    qs = saved["questions"]
    assert [q["id"] for q in qs] == ["q1", "q2", "q3", "q4", "q5"] and [o["id"] for o in qs[0]["options"]] == ["o1", "o2", "o3"] and saved["slug"] == "test-survey"
    again = core.save_survey(s, "o", {**saved, "questions": [{**qs[0], "text": "Do you agree (corrected)?"}] + qs[1:]}, saved["id"])
    assert again["questions"][0]["id"] == "q1" and again["questions"][0]["options"][1]["id"] == "o2"                  # a corrected wording keeps every id
    for t in core.template_list():
        draft = core.from_template(t["key"])
        saved_t = core.save_survey(s, "o", {**draft, "slug": t["key"], "status": "draft"})
        assert saved_t["questions"] and saved_t["status"] == "draft", t["key"]
    assert {t["key"] for t in core.template_list()} >= {"quick-poll", "political-poll", "brand-tracker", "customer-satisfaction", "team-pulse", "media-habits"}
    demo = core.demographics()
    assert len(demo) == 4 and len(demo[3]["options"]) == len(core.INDIA_STATES) == 36


def test_what_a_person_must_fix_is_said_plainly():
    s, _ = make()
    with pytest.raises(ValueError, match="title cannot be empty"):
        core.save_survey(s, "o", base(title=" "))
    with pytest.raises(ValueError, match="at least two options"):
        core.save_survey(s, "o", base(questions=[{"type": "single", "text": "One option", "options": opts("Only")}]))
    with pytest.raises(ValueError, match="comes before"):
        core.save_survey(s, "o", base(questions=[{"type": "single", "text": "A", "options": opts("x", "y"), "show_if": {"q": "q2", "op": "is", "values": ["o1"]}}, {"type": "text", "text": "B"}]))
    with pytest.raises(ValueError, match="add at least one question"):
        core.save_survey(s, "o", {"title": "Empty", "status": "live", "questions": [{"type": "section", "text": "Only a heading"}]})
    with pytest.raises(ValueError, match="consent"):
        core.save_survey(s, "o", base(title="Consent test", contact_mode="optional", contact_fields=["mobile"]))
    with pytest.raises(ValueError, match="choose which contact"):
        core.save_survey(s, "o", base(title="Fields test", contact_mode="optional", consent_text="I agree"))
    with pytest.raises(ValueError, match="after the opening"):
        core.save_survey(s, "o", base(title="Dates test", opens_at="2026-11-02T10:00:00+00:00", closes_at="2026-11-01T10:00:00+00:00"))
    with pytest.raises(ValueError, match="https://"):
        core.save_survey(s, "o", base(title="Redirect test", redirect_url="http://x.com"))
    with pytest.raises(ValueError, match="at most three"):
        core.save_survey(s, "o", base(title="Poll test", kind="poll"))
    live(s)
    with pytest.raises(ValueError, match="already used"):
        core.save_survey(s, "o", base())


def test_once_answers_exist_nothing_that_holds_an_answer_can_be_deleted_or_retyped_but_it_can_be_archived():
    s, _ = make()
    saved = live(s)
    row = core.get_survey(s, saved["id"])
    answer(s, row, {"q1": {"o": "o1"}})
    qs = saved["questions"]
    with pytest.raises(ValueError, match="cannot be deleted"):
        core.save_survey(s, "o", {**saved, "questions": qs[1:]}, saved["id"])
    with pytest.raises(ValueError, match="type"):
        core.save_survey(s, "o", {**saved, "questions": [{**qs[0], "type": "multi"}] + qs[1:]}, saved["id"])
    with pytest.raises(ValueError, match="option"):
        core.save_survey(s, "o", {**saved, "questions": [{**qs[0], "options": qs[0]["options"][:2]}] + qs[1:]}, saved["id"])
    ok = core.save_survey(s, "o", {**saved, "questions": [{**qs[0], "options": qs[0]["options"][:2] + [{**qs[0]["options"][2], "archived": True}]}] + [{**qs[4], "archived": True}] + qs[1:4]}, saved["id"])
    assert ok["questions"][0]["options"][2]["archived"] is True
    pub = core.public_definition(core.get_survey(s, saved["id"]))
    assert [o["id"] for o in pub["questions"][0]["options"]] == ["o1", "o2"] and "q5" not in [q["id"] for q in pub["questions"]]       # hidden from new respondents
    assert core.survey_results(s, row)["questions"][0]["options"][2]["id"] == "o3"                                            # still in the results


# ------------------------------------------------------------------------------------------------------------------------------ answers
def test_every_question_type_is_checked_on_the_server():
    s, _ = make()
    saved = live(s, questions=[
        {"type": "single", "text": "S", "required": True, "options": opts("a", "b") + [{"text": "Other", "other": True}]},
        {"type": "multi", "text": "M", "max_select": 2, "options": opts("a", "b", "c") + [{"text": "None", "exclusive": True}]},
        {"type": "dropdown", "text": "D", "options": opts("a", "b")},
        {"type": "rating", "text": "R", "scale_max": 5}, {"type": "nps", "text": "N"}, {"type": "number", "text": "Num", "min": 0, "max": 100},
        {"type": "text", "text": "T", "max_length": 20}, {"type": "ranking", "text": "K", "options": opts("a", "b", "c")}])
    qs = saved["questions"]
    ok = core.validate_answers(qs, {"q1": {"o": "o1"}, "q2": {"o": ["o1", "o2"]}, "q3": {"o": "o2"}, "q4": {"n": 4}, "q5": {"n": 10}, "q6": {"n": 55.5}, "q7": {"t": " hello "}, "q8": {"o": ["o3", "o1", "o2"]}, "junk": {"o": "x"}})
    assert ok["q1"] == {"o": "o1"} and ok["q6"] == {"n": 55.5} and ok["q7"] == {"t": "hello"} and "junk" not in ok
    assert core.validate_answers(qs, {"q1": {"o": "o3", "t": "my own"}})["q1"] == {"o": "o3", "t": "my own"}
    cases = [({}, "q1"), ({"q1": {"o": "o9"}}, "q1"), ({"q1": {"o": "o3"}}, "q1"), ({"q1": {"o": "o1"}, "q2": {"o": ["o1", "o2", "o3"]}}, "q2"), ({"q1": {"o": "o1"}, "q2": {"o": ["o1", "o4"]}}, "q2"),
             ({"q1": {"o": "o1"}, "q4": {"n": 6}}, "q4"), ({"q1": {"o": "o1"}, "q4": {"n": 2.5}}, "q4"), ({"q1": {"o": "o1"}, "q5": {"n": 11}}, "q5"), ({"q1": {"o": "o1"}, "q6": {"n": 101}}, "q6"),
             ({"q1": {"o": "o1"}, "q6": {"n": float("nan")}}, "q6"), ({"q1": {"o": "o1"}, "q7": {"t": "x" * 21}}, "q7"), ({"q1": {"o": "o1"}, "q8": {"o": ["o1", "o2"]}}, "q8"), ({"q1": "text"}, "q1")]
    for answers, bad in cases:
        with pytest.raises(core.AnswerError) as e:
            core.validate_answers(qs, answers)
        assert bad in e.value.errors, (answers, e.value.errors)
    with pytest.raises(core.AnswerError):
        core.validate_answers(qs, "not an object")


def test_a_question_that_depends_on_another_is_asked_only_when_it_should_be():
    s, _ = make()
    saved = live(s, questions=[{"type": "single", "text": "Do you own one?", "options": opts("Yes", "No")},
                               {"type": "text", "text": "Which one?", "required": True, "show_if": {"q": "q1", "op": "is", "values": ["o1"]}},
                               {"type": "text", "text": "Why not?", "required": True, "show_if": {"q": "q1", "op": "not", "values": ["o1"]}}])
    qs = saved["questions"]
    assert core.validate_answers(qs, {"q1": {"o": "o2"}, "q3": {"t": "price"}, "q2": {"t": "ignored"}}) == {"q1": {"o": "o2"}, "q3": {"t": "price"}}     # the answer to a hidden question is dropped
    with pytest.raises(core.AnswerError) as e:
        core.validate_answers(qs, {"q1": {"o": "o1"}})
    assert list(e.value.errors) == ["q2"]


def test_a_response_is_stored_anonymously_and_the_rules_of_the_survey_are_enforced():
    s, _ = make()
    saved = live(s, max_responses=2)
    row = core.get_survey(s, saved["id"])
    r = answer(s, row, {"q1": {"o": "o1"}}, device="1.2.3.4|Mozilla")
    stored = s.get(SurveyResponse, r["id"])
    assert stored.contact_json is None and "1.2.3.4" not in (stored.device_hash or "") and len(stored.device_hash) == 64 and "1.2.3.4" not in json.dumps(stored.__dict__, default=str)
    token = "t" * 20
    answer(s, row, {"q1": {"o": "o2"}}, token=token)
    with pytest.raises(ValueError, match="closed"):
        answer(s, row, {"q1": {"o": "o1"}})                                                                      # full: two answers allowed
    s.query(Survey).filter_by(id=row.id).update({"max_responses": None})
    s.commit()
    with pytest.raises(ValueError, match="already answered"):
        answer(s, row, {"q1": {"o": "o1"}}, token=token)                                                         # the same browser twice
    row.allow_repeat = True
    answer(s, row, {"q1": {"o": "o1"}}, token=token)                                                              # a kiosk may repeat
    with pytest.raises(core.AnswerError):
        answer(s, row, {})


def test_closed_scheduled_expired_and_draft_surveys_take_no_answers():
    s, _ = make()
    for name, over, text in (("closed", {"status": "closed"}, "closed"), ("future", {"opens_at": (NOW + timedelta(days=1)).isoformat()}, "not opened"), ("past", {"closes_at": (NOW - timedelta(days=1)).isoformat()}, "closed"),
                             ("draft", {"status": "draft"}, "closed")):
        row = core.get_survey(s, core.save_survey(s, "o", base(title=name, **over))["id"])
        with pytest.raises(ValueError, match=text):
            answer(s, row, {"q1": {"o": "o1"}})
    assert core.state_of(core.get_survey(s, core.save_survey(s, "o", base(title="open one"))["id"]), NOW) == "open"


def test_devices_and_very_fast_answers_are_flagged_not_lost():
    s, _ = make()
    row = core.get_survey(s, live(s)["id"])
    answer(s, row, {"q1": {"o": "o1"}}, device="9.9.9.9|UA")
    r2 = answer(s, row, {"q1": {"o": "o2"}}, device="9.9.9.9|UA")
    assert s.get(SurveyResponse, r2["id"]).status == "flagged" and "same device" in s.get(SurveyResponse, r2["id"]).flags
    r3 = answer(s, row, {"q1": {"o": "o2"}, "q3": {"n": 3}, "q4": {"n": 8}}, started_ms=(NOW - timedelta(seconds=1)).timestamp() * 1000)
    assert "very fast" in s.get(SurveyResponse, r3["id"]).flags
    assert core.response_count(s, row.id) == 3 and core.response_count(s, row.id, include_flagged=False) == 1
    core.set_response_status(s, r2["id"], "removed")
    assert core.response_count(s, row.id) == 2 and core.purge_removed(s, row.id) == 1


def test_contact_details_need_consent_and_are_kept_apart():
    s, _ = make()
    row = core.get_survey(s, live(s, title="With contact", contact_mode="required", contact_fields=["name", "mobile"], consent_text="I agree to be called.")["id"])
    with pytest.raises(core.AnswerError, match="details"):
        answer(s, row, {"q1": {"o": "o1"}}, contact={"name": "Asha"}, consent=True)
    with pytest.raises(core.AnswerError, match="mobile"):
        answer(s, row, {"q1": {"o": "o1"}}, contact={"name": "Asha", "mobile": "12"}, consent=True)
    with pytest.raises(core.AnswerError, match="agree"):
        answer(s, row, {"q1": {"o": "o1"}}, contact={"name": "Asha", "mobile": "98765 43210"}, consent=False)
    r = answer(s, row, {"q1": {"o": "o1"}}, contact={"name": "Asha", "mobile": "98765 43210", "email": "ignored@x.in"}, consent=True)
    assert json.loads(s.get(SurveyResponse, r["id"]).contact_json) == {"name": "Asha", "mobile": "9876543210"}      # only the asked fields, digits only
    assert "Asha" not in json.dumps(core.survey_results(s, row)) and "9876543210" not in core.export_csv(s, row).decode()
    assert "9876543210" in core.export_csv(s, row, include_contact=True).decode()


def test_invitation_only_links_work_once_and_only_their_hash_is_kept():
    s, _ = make()
    row = core.get_survey(s, live(s, title="Invited", access="invite")["id"])
    with pytest.raises(ValueError, match="invitation"):
        answer(s, row, {"q1": {"o": "o1"}})
    made = core.create_invites(s, row, ["Asha", "Ravi"])
    assert [m["label"] for m in made] == ["Asha", "Ravi"] and len({m["code"] for m in made}) == 2
    assert not [c for c in made if c["code"] in json.dumps([i.code_hash for i in s.query(__import__("collector.storage", fromlist=["SurveyInvite"]).SurveyInvite).all()])]
    answer(s, row, {"q1": {"o": "o1"}}, invite=made[0]["code"])
    with pytest.raises(ValueError, match="already answered"):
        answer(s, row, {"q1": {"o": "o1"}}, invite=made[0]["code"])
    with pytest.raises(ValueError, match="invitation"):
        answer(s, row, {"q1": {"o": "o1"}}, invite="made-up-code")
    assert core.invite_summary(s, row.id)["used"] == 1 and core.invite_summary(s, row.id)["total"] == 2
    assert len(core.create_invites(s, row, count=3)) == 3
    with pytest.raises(ValueError, match="between 1 and 5000"):
        core.create_invites(s, row)


# ------------------------------------------------------------------------------------------------------------------------------ results
def filled(s, sampling="self_selected", population=None, n=40):
    saved = live(s, title="Results check", sampling=sampling, population_size=population)
    row = core.get_survey(s, saved["id"])
    for i in range(n):
        a = {"q1": {"o": "o1" if i % 4 else "o2"}, "q2": {"o": ["o1", "o2"] if i % 2 else ["o3"]}, "q3": {"n": 1 + i % 5}, "q4": {"n": 10 if i % 3 == 0 else (8 if i % 3 == 1 else 3)}}
        if i % 10 == 0:
            a["q5"] = {"t": f"comment {i}"}
        when = NOW - timedelta(hours=i)
        answer(s, row, a, started_ms=(when - timedelta(seconds=120)).timestamp() * 1000, now=when)
    return row


def test_results_count_percent_and_say_what_they_are():
    s, _ = make()
    row = filled(s)
    res = core.survey_results(s, row)
    q1 = res["questions"][0]
    assert res["totals"]["counted"] == 40 and q1["n"] == 40 and {o["id"]: o["count"] for o in q1["options"]} == {"o1": 30, "o2": 10, "o3": 0} and q1["options"][0]["pct"] == 75.0
    assert "moe" not in q1["options"][0] and "no margin of error" in res["note"] and res["scientific"] is False                      # self-selected: no margin of error is claimed
    q3, q4, q5 = res["questions"][2], res["questions"][3], res["questions"][4]
    assert q3["mean"] == 3.0 and q3["median"] == 3 and len(q3["distribution"]) == 5 and sum(d["count"] for d in q3["distribution"]) == 40
    assert q4["nps"] == {"promoters": 14, "passives": 13, "detractors": 13, "score": 2.5} and len(q4["distribution"]) == 11
    assert len(q5["texts"]) == 4 and res["by_day"] and res["totals"]["median_seconds"] == 120


def test_a_random_sample_gets_a_margin_of_error_that_shrinks_with_more_answers():
    s, _ = make()
    row = filled(s, sampling="random_sample", population=100_000)
    opt = core.survey_results(s, row)["questions"][0]["options"][0]
    assert opt["moe"] == pytest.approx(core.margin_of_error(0.75, 40, 100_000)) and 12 < opt["moe"] < 14
    assert core.margin_of_error(0.5, 1000) == 3.1 and core.margin_of_error(0.5, 400) > core.margin_of_error(0.5, 1600) and core.margin_of_error(0.5, 1) == 100.0
    assert core.margin_of_error(0.5, 900, 1000) < core.margin_of_error(0.5, 900, 1_000_000)                                       # a small known population tightens it
    assert "Random sample" in core.survey_results(s, row)["note"]


def test_flagged_and_removed_answers_can_be_left_out_and_a_split_by_group_is_shown():
    s, _ = make()
    saved = live(s, title="Split", questions=[{"type": "single", "text": "Age", "options": opts("Young", "Old")}, {"type": "single", "text": "Vote", "options": opts("A", "B")}])
    row = core.get_survey(s, saved["id"])
    for age, vote, n in (("o1", "o1", 6), ("o1", "o2", 2), ("o2", "o1", 1), ("o2", "o2", 7)):
        for _ in range(n):
            answer(s, row, {"q1": {"o": age}, "q2": {"o": vote}})
    x = core.crosstab(s, row, "q2", "q1")
    assert [c["n"] for c in x["columns"]] == [8, 8] and x["rows"][0]["cells"] == [{"count": 6, "pct": 75.0}, {"count": 1, "pct": 12.5}] and x["rows"][1]["cells"][1]["pct"] == 87.5
    with pytest.raises(ValueError, match="choose a choice question"):
        core.crosstab(s, row, "q9", "q1")
    one = s.query(SurveyResponse).first()
    core.set_response_status(s, one.id, "removed")
    assert core.survey_results(s, row)["totals"]["counted"] == 15 and core.survey_results(s, row)["totals"]["removed"] == 1
    s.query(SurveyResponse).filter(SurveyResponse.id == one.id + 1).update({"status": "flagged"})
    s.commit()
    assert core.survey_results(s, row, include_flagged=False)["totals"]["counted"] == 14


def test_public_results_follow_the_settings_and_never_show_text_or_splits():
    s, _ = make()
    row = filled(s, n=12)
    assert core.public_results(row, s) is None                                                                       # private by default
    row.results_mode = "public"
    row.min_results = 50
    assert core.public_results(row, s) == {"waiting": True, "needed": 50, "have": 12}
    row.min_results = 10
    res = core.public_results(row, s)
    assert res["n"] == 12 and [q["type"] for q in res["questions"]] == ["single", "multi", "rating", "nps"] and "texts" not in json.dumps(res) and "other_texts" not in json.dumps(res)
    row.results_hold = True
    assert core.public_results(row, s) is None                                                                       # a hold overrides everything
    row.results_hold, row.results_mode = False, "after_close"
    assert core.public_results(row, s) is None
    row.status = "closed"
    assert core.public_results(row, s)["n"] == 12


def test_exports_are_complete_safe_in_a_spreadsheet_and_open():
    s, _ = make()
    saved = live(s, title="Export", questions=[{"type": "single", "text": "Pick", "options": [{"text": "=SUM(1+1)"}, {"text": "Fine"}]}, {"type": "text", "text": "Say"}])
    row = core.get_survey(s, saved["id"])
    answer(s, row, {"q1": {"o": "o1"}, "q2": {"t": "+cmd|' /C calc'!A0"}})
    data = core.export_csv(s, row)
    assert data.startswith("\ufeff".encode("utf-8")) and b"q1: Pick" in data
    text = data.decode("utf-8-sig")
    assert "\n =SUM(1+1)" in text.replace("\r", "").replace(",", "\n") or ' =SUM(1+1)' in text and ' +cmd' in text                  # a formula is kept as text
    wb = load_workbook(io.BytesIO(core.export_xlsx(s, row)))
    assert wb.sheetnames == ["Answers", "Summary"] and wb["Answers"].max_row == 2 and wb["Summary"]["B1"].value == "Export"


# ------------------------------------------------------------------------------------------------------------------------------ the routes and pages
def user(s, email, admin=False):
    u = User(email=email, password_hash=hash_password("a long enough pass"), plan="sx_enterprise", requested_plan="sx_enterprise", is_admin=admin, active=True, approval_status="approved", full_name=email, organization="Org")
    s.add(u)
    s.commit()
    return u


def admin_call(engine):
    app = survey_admin_application(lambda: Session(engine))

    def call(who, method, route, body=None, query=""):
        raw = json.dumps(body).encode() if body is not None else b""
        out = {}
        env = {"REQUEST_METHOD": method, "PATH_INFO": "/api/v1/surveys/" + route, "QUERY_STRING": query, "wsgi.input": io.BytesIO(raw), "CONTENT_LENGTH": str(len(raw)),
               "HTTP_AUTHORIZATION": ("Bearer " + issue_token(AuthIdentity(who.id, who.email, who.plan, who.is_admin))) if who else ""}
        res = app(env, lambda status, headers: out.update(status=status, headers=dict(headers)))
        data = b"".join(res)
        ctype = out["headers"].get("Content-Type", "")
        return int(out["status"].split()[0]), (json.loads(data) if ctype.startswith("application/json") else data), out["headers"]
    return call


def public_call(engine):
    app = survey_public_application(lambda: Session(engine))

    def call(method, slug, action="", body=None, ip="5.5.5.5"):
        raw = json.dumps(body).encode() if body is not None else b""
        out = {}
        env = {"REQUEST_METHOD": method, "PATH_INFO": f"/api/v1/public/survey/{slug}" + (f"/{action}" if action else ""), "QUERY_STRING": "", "wsgi.input": io.BytesIO(raw), "CONTENT_LENGTH": str(len(raw)),
               "REMOTE_ADDR": ip, "HTTP_USER_AGENT": "pytest"}
        res = app(env, lambda status, headers: out.update(status=status))
        return int(out["status"].split()[0]), json.loads(b"".join(res))
    return call


def test_only_an_administrator_runs_surveys_and_the_whole_flow_works_over_http():
    s, engine = make()
    admin, member = user(s, "admin@x.in", True), user(s, "member@x.in")
    call = admin_call(engine)
    assert call(None, "GET", "list")[0] == 401 and call(member, "GET", "list")[0] == 403 and call(member, "POST", "save", base())[0] == 403
    code, body, _ = call(admin, "GET", "list")
    assert code == 200 and body["surveys"] == [] and len(body["templates"]) == 6
    code, body, _ = call(admin, "GET", "template", query="key=political-poll")
    assert code == 200 and body["survey"]["category"] == "political" and len(body["survey"]["questions"]) == 8
    code, body, _ = call(admin, "POST", "save", {**body["survey"], "slug": "city-poll", "status": "draft"})
    sid = body["survey"]["id"]
    assert code == 200 and body["survey"]["state"] == "draft"
    assert call(admin, "POST", "save", base(title="Other", slug="city-poll"))[0] == 400
    code, body, _ = call(admin, "POST", "preview", {"id": sid, "title": "City poll", "questions": call(admin, "GET", "get", query=f"id={sid}")[1]["survey"]["questions"]})
    assert code == 200 and "noindex" in body["html"] and 'id="def"' in body["html"]
    assert call(admin, "POST", "status", {"id": sid, "status": "live"})[1]["survey"]["state"] == "open"
    pc = public_call(engine)
    code, body = pc("GET", "city-poll")
    assert code == 200 and body["accepting"] is True and body["survey"]["questions"][0]["text"].startswith("If an election") and "closed" not in json.dumps(body["survey"]).lower()[:0]
    assert pc("GET", "city-poll", "results")[0] == 403 and pc("GET", "nope")[0] == 404
    code, body = pc("POST", "city-poll", "respond", {"answers": {"q1": {"o": "o1"}, "q3": {"n": 4}}, "token": "x" * 24, "started_ms": 1})
    assert code == 200 and body == {"ok": True}
    code, body = pc("POST", "city-poll", "respond", {"answers": {"q1": {"o": "o1"}}, "token": "x" * 24, "started_ms": 1})
    assert code == 400 and "already" in body["error"]
    code, body = pc("POST", "city-poll", "respond", {"answers": {"q1": {"o": "nope"}}, "token": "y" * 24})
    assert code == 400 and "q1" in body["errors"]
    assert call(admin, "GET", "results", query=f"id={sid}")[1]["totals"]["counted"] == 1
    assert call(admin, "GET", "responses", query=f"id={sid}")[1]["responses"][0]["answers"]["q1"] == "Party or candidate A"
    csv_code, csv_body, headers = call(admin, "GET", "export", query=f"id={sid}&format=csv")
    assert csv_code == 200 and headers["Content-Disposition"].endswith('city-poll.csv"') and csv_body.startswith(b"\xef\xbb\xbf")
    assert call(admin, "GET", "export", query=f"id={sid}&format=xlsx")[2]["Content-Disposition"].endswith('city-poll.xlsx"')
    assert call(admin, "POST", "delete", {"id": sid})[0] == 400                                                       # it has an answer: it cannot be deleted by mistake
    assert call(admin, "POST", "status", {"id": sid, "status": "closed"})[1]["survey"]["state"] == "closed"
    assert pc("POST", "city-poll", "respond", {"answers": {"q1": {"o": "o1"}}, "token": "z" * 24})[0] == 400
    assert call(admin, "POST", "duplicate", {"id": sid})[1]["survey"]["slug"] == "city-poll-copy"
    assert call(admin, "GET", "get", query="id=9999")[0] == 404


def test_invitation_links_over_http_and_the_hidden_trap_field():
    s, engine = make()
    admin = user(s, "admin@x.in", True)
    call, pc = admin_call(engine), public_call(engine)
    sid = call(admin, "POST", "save", base(title="Private panel", access="invite"))[1]["survey"]["id"]
    code, body, _ = call(admin, "POST", "invites/create", {"id": sid, "labels": ["Asha"]})
    link = body["links"][0]["url"]
    assert code == 200 and "?i=" in link and "shown only once" in body["message"]
    assert pc("POST", "private-panel", "respond", {"answers": {"q1": {"o": "o1"}}, "token": "a" * 24})[0] == 400
    assert pc("POST", "private-panel", "respond", {"answers": {"q1": {"o": "o1"}}, "token": "a" * 24, "invite": link.split("?i=")[1]})[0] == 200
    open_id = call(admin, "POST", "save", base(title="Open one"))[1]["survey"]["id"]
    assert call(admin, "POST", "invites/create", {"id": open_id, "count": 2})[0] == 400
    assert pc("POST", "open-one", "respond", {"answers": {"q1": {"o": "o1"}}, "token": "b" * 24, "hp": "i am a bot"}) == (200, {"ok": True})
    assert core.response_count(s, open_id) == 0                                                                        # the bot was told ok and nothing was stored


def test_answers_are_rate_limited_by_connection():
    s, engine = make()
    admin = user(s, "admin@x.in", True)
    admin_call(engine)(admin, "POST", "save", base(title="Busy"))
    pc = public_call(engine)
    codes = [pc("POST", "busy", "respond", {"answers": {"q1": {"o": "o1"}}, "token": f"{i:024d}"}, ip="7.7.7.7")[0] for i in range(10)]
    assert codes[:8] == [200] * 8 and codes[8:] == [429, 429]
    assert pc("POST", "busy", "respond", {"answers": {"q1": {"o": "o1"}}, "token": "9" * 24}, ip="8.8.8.8")[0] == 200


def page(engine, path):
    out = {}
    res = survey_site_response(lambda: Session(engine), {"REQUEST_METHOD": "GET", "PATH_INFO": path}, lambda status, headers: out.update(status=status, headers=dict(headers)), [])
    return (None, None, "") if res is None else (int(out["status"].split()[0]), out["headers"], b"".join(res).decode("utf-8"))


def test_the_pages_are_safe_honest_and_indexed_only_when_listed():
    s, engine = make()
    admin = user(s, "admin@x.in", True)
    call = admin_call(engine)
    evil = '</script><script>alert(1)</script><img src=x onerror=alert(2)>'
    call(admin, "POST", "save", base(title="Open poll " + evil, slug="open-poll", description=evil, listed=True, sponsor="Paid by Example Ltd"))
    call(admin, "POST", "save", base(title="Hidden one", slug="hidden-one"))
    call(admin, "POST", "save", base(title="A draft", slug="a-draft", status="draft"))
    code, headers, html = page(engine, "/survey/open-poll")
    assert code == 200 and 'content="index,follow"' in html and "alert(1)</script>" not in html.replace("\\u003c/script\\u003e", "") and "<img src=x" not in html and 'id="def"' in html and "Paid by Example Ltd" in html
    assert json.loads(re.search(r'<script type="application/json" id="def">(.*?)</script>', html, re.S).group(1))["questions"][0]["text"] == "Do you agree?"
    assert 'content="noindex,nofollow"' in page(engine, "/survey/hidden-one")[2]                                      # not listed: kept out of search
    assert page(engine, "/survey/a-draft")[0] == 404 and page(engine, "/survey/nothing")[0] == 404 and page(engine, "/platform") == (None, None, "")
    code, _, html = page(engine, "/polls")
    assert code == 200 and "open-poll" in html and "hidden-one" not in html and "a-draft" not in html
    assert "The results of this poll are not public" in page(engine, "/survey/open-poll/results")[2]
    from api import services_cms
    sm = services_cms.sitemap_xml(s)
    assert "/survey/open-poll" in sm and "/polls" in sm and "hidden-one" not in sm
    call(admin, "POST", "save", base(title="Open poll", slug="open-poll", listed=True, results_mode="public", min_results=1, id=1, method_note="Sampled by phone."))
    pc = public_call(engine)
    pc("POST", "open-poll", "respond", {"answers": {"q1": {"o": "o1"}, "q5": {"t": "secret comment"}}, "token": "k" * 24})
    results = page(engine, "/survey/open-poll/results")[2]
    assert "About these results" in results and "self-selected" in results.lower() and "Sampled by phone." in results and "secret comment" not in results
    assert pc("GET", "open-poll", "results")[1]["n"] == 1 and "secret comment" not in json.dumps(pc("GET", "open-poll", "results")[1])


def test_the_scripts_of_the_pages_parse():
    node = shutil.which("node")
    if not node:
        pytest.skip("node is not installed")
    with tempfile.TemporaryDirectory() as d:
        for name, text in (("form.js", survey_pages.JS), ("common.js", "(()=>{" + survey_pages.JS_COMMON + "})();")):
            path = Path(d) / name
            path.write_text(text, encoding="utf-8")
            done = subprocess.run([node, "--check", str(path)], capture_output=True, text=True)
            assert done.returncode == 0, (name, done.stderr[:400])


def test_the_web_app_serves_the_poll_pages_and_keeps_the_admin_api_private():
    import web
    out = {}

    def get(path):
        res = web.application({"REQUEST_METHOD": "GET", "PATH_INFO": path, "QUERY_STRING": "", "REMOTE_ADDR": "127.0.0.1"}, lambda status, headers: out.update(status=status, headers=dict(headers)))
        return out["status"], b"".join(res)
    assert get("/polls")[0].startswith("200") and get("/survey/not-a-survey")[0].startswith("404")
    assert get("/api/v1/surveys/list")[0].startswith("401") and get("/api/v1/public/survey/not-a-survey")[0].startswith("404")
