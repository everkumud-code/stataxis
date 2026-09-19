import json
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from api.auth import AuthIdentity
from api.auth_guard import protect_application
from api.http import wsgi_application
from collector.storage import Base, Channel, Observation, Video


def make_app():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)

    with Session(engine) as session:
        channel = Channel(
            youtube_channel_id="channel-1",
            name="Configured Channel",
            network="Network",
            language="Hindi",
            region="North India",
        )
        session.add(channel)
        session.flush()
        video = Video(
            youtube_video_id="video1234567",
            channel_id=channel.id,
            title="Export me",
        )
        session.add(video)
        session.flush()
        session.add(
            Observation(
                video_id=video.id,
                channel_id=channel.id,
                observed_at=datetime(2026, 9, 1, 12, tzinfo=UTC),
                view_count=100,
                like_count=5,
                comment_count=2,
                concurrent_viewers=None,
                is_live=False,
                classification="VOD",
            )
        )
        session.commit()

    return protect_application(wsgi_application(lambda: Session(engine)))


def call(app, environ):
    captured = {}
    body = app(
        environ,
        lambda status, headers: captured.update(status=status, headers=headers),
    )
    return captured, b"".join(body)


def test_report_export_requires_authentication():
    captured, body = call(
        make_app(),
        {
            "REQUEST_METHOD": "GET",
            "PATH_INFO": "/api/v1/reports/export",
            "QUERY_STRING": "",
        },
    )
    assert captured["status"] == "401 Unauthorized"
    assert json.loads(body) == {"error": "authentication required"}


def test_report_export_rejects_free_plan(monkeypatch):
    monkeypatch.setattr(
        "api.auth_guard.authenticate",
        lambda _authorization: AuthIdentity(1, "free@example.org", "sx_free"),
    )
    captured, body = call(
        make_app(),
        {
            "REQUEST_METHOD": "GET",
            "PATH_INFO": "/api/v1/reports/export",
            "QUERY_STRING": "language=Hindi&region=North%20India&classification=VOD",
            "HTTP_AUTHORIZATION": "Bearer test",
        },
    )
    assert captured["status"] == "403 Forbidden"
    assert json.loads(body) == {"error": "premium access required"}


def test_report_export_returns_xlsx_for_premium_plan(monkeypatch):
    monkeypatch.setattr(
        "api.auth_guard.authenticate",
        lambda _authorization: AuthIdentity(2, "pro@example.org", "sx_pro"),
    )
    captured, body = call(
        make_app(),
        {
            "REQUEST_METHOD": "GET",
            "PATH_INFO": "/api/v1/reports/export",
            "QUERY_STRING": "language=Hindi&region=North%20India&classification=VOD",
            "HTTP_AUTHORIZATION": "Bearer test",
        },
    )
    assert captured["status"] == "200 OK"
    headers = dict(captured["headers"])
    assert headers["Content-Type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert headers["Content-Disposition"] == "attachment; filename=stataxis-report.xlsx"
    workbook = load_workbook(BytesIO(body), read_only=True)
    assert workbook.sheetnames == ["StatAxis Data", "StatAxis Intelligence"]
    assert workbook["StatAxis Data"].max_row == 2
