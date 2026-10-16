"""Facebook, Instagram and X: the connectors (against invented platform answers), reading and storing, sheets, the rankings and growth, the HTTP routes and who may use them, and how long figures are kept.
Invented data in a scratch database; nothing here touches a platform."""
import io
import json
import urllib.error
from datetime import UTC, datetime, timedelta

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from api import social
from api.auth import AuthIdentity, hash_password, issue_token
from api.social_http import social_application
from collector.social import service
from collector.social import settings as social_settings
from collector.social.base import SocialError, clean_handle, http_json, parse_time, to_int, valid_handle
from collector.social.connectors import Credentials, Facebook, Instagram, X, connectors
from collector.social.retention import run_social_retention
from collector.storage import AppSetting, Base, Channel, ChannelStats, SocialAccount, SocialPost, SocialPostStats, SocialStats, User

T0 = datetime(2026, 10, 16, 6, 0, tzinfo=UTC)


@pytest.fixture(autouse=True)
def secret(monkeypatch):
    monkeypatch.setenv("STAXIS_AUTH_SECRET", "test-secret")
    for name in ("STAXIS_META_TOKEN", "STAXIS_META_IG_USER_ID", "STAXIS_X_BEARER_TOKEN"):
        monkeypatch.delenv(name, raising=False)


def make_session():
    engine = create_engine("sqlite://", future=True, poolclass=StaticPool, connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    return Session(engine), engine


def channel(s, key, name, segment="news", language="Hindi", handle=None):
    c = Channel(youtube_channel_id=key, name=name, network="n", language=language, region="India", segment=segment, handle=handle)
    s.add(c)
    s.flush()
    return c


class FakeResponse:
    def __init__(self, body):
        self.body = body

    def read(self):
        return json.dumps(self.body).encode()

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False


def opener(routes):
    """routes: list of (url fragment, body or an HTTP error code). The first fragment found in the address answers."""
    calls = []

    def _open(req, timeout):
        calls.append(req.full_url)
        for fragment, body in routes:
            if fragment in req.full_url:
                if isinstance(body, int):
                    raise urllib.error.HTTPError(req.full_url, body, "err", {}, io.BytesIO(json.dumps({"error": {"message": "x", "code": body}}).encode()))
                return FakeResponse(body)
        raise urllib.error.HTTPError(req.full_url, 404, "nf", {}, io.BytesIO(b"{}"))
    _open.calls = calls
    return _open


KEYS = Credentials(meta_token="T" * 30, meta_ig_user_id="1789", x_bearer="B" * 30)


# ------------------------------------------------------------------------------------------------------------------------------ helpers
def test_handles_are_cleaned_and_checked():
    assert clean_handle("https://x.com/stataxis/", "x") == "stataxis"
    assert clean_handle("@StatAxis", "x") == "StatAxis"
    assert clean_handle("https://www.instagram.com/aajtak/?hl=en", "instagram") == "aajtak"
    assert clean_handle("https://www.facebook.com/pages/Name/12345", "facebook") == "12345"
    assert valid_handle("stataxis", "x") and not valid_handle("a-very-long-handle-over-15", "x") and not valid_handle("bad handle", "instagram")
    assert to_int("1,234") == 1234 and to_int("-5") is None and to_int("") is None and to_int("abc") is None
    assert parse_time("2026-10-12T09:30:00+0000") == datetime(2026, 10, 12, 9, 30, tzinfo=UTC) and parse_time("2026-10-12T09:30:00Z") == datetime(2026, 10, 12, 9, 30, tzinfo=UTC) and parse_time("nope") is None


def test_platform_refusals_become_clear_errors_without_the_key():
    for status, code in ((429, "rate_limited"), (401, "auth"), (403, "unsupported"), (404, "not_found"), (500, "error")):
        with pytest.raises(SocialError) as err:
            http_json("https://graph.facebook.com/x?access_token=SECRETKEY", opener=opener([("graph", status)]))
        assert err.value.code == code and "SECRETKEY" not in str(err.value)


# ------------------------------------------------------------------------------------------------------------------------------ connectors
def test_instagram_business_discovery_is_read():
    body = {"business_discovery": {"id": "55", "username": "aajtak", "name": "Aaj Tak", "followers_count": 12_000_000, "follows_count": 90, "media_count": 8000,
                                   "media": {"data": [{"id": "m1", "caption": "Breaking:  big news", "media_type": "IMAGE", "permalink": "https://instagram.com/p/m1", "timestamp": "2026-10-15T10:00:00+0000", "like_count": 500, "comments_count": 20},
                                                      {"id": "m2", "timestamp": "2026-10-14T10:00:00+0000"}]}}, "id": "1789"}
    op = opener([("graph.facebook.com/v21.0/1789", body)])
    got = Instagram(KEYS, op).fetch("aajtak")
    assert got.account.followers == 12_000_000 and got.account.following == 90 and got.account.posts_total == 8000 and got.account.url.endswith("/aajtak/")
    assert [p.post_id for p in got.posts] == ["m1", "m2"] and got.posts[0].likes == 500 and got.posts[0].comments == 20 and got.posts[0].text == "Breaking: big news"
    assert got.posts[1].likes is None and got.posts[1].comments is None                      # not given = empty, never zero
    assert "business_discovery.username(aajtak)" in op.calls[0].replace("%28", "(").replace("%29", ")")


def test_instagram_personal_account_and_missing_keys_are_explained():
    with pytest.raises(SocialError) as err:
        Instagram(KEYS, opener([("graph", {"id": "1789"})])).fetch("someone")
    assert err.value.code == "unsupported" and "business" in str(err.value)
    with pytest.raises(SocialError) as err:
        Instagram(Credentials(), opener([])).fetch("someone")
    assert err.value.code == "not_configured"
    assert not Instagram(Credentials(meta_token="x" * 30), None).ready() and Instagram(KEYS, None).ready()


def test_facebook_page_reactions_comments_and_shares():
    page = {"id": "99", "name": "Aaj Tak", "followers_count": 5_000_000, "fan_count": 4_900_000, "link": "https://facebook.com/aajtak"}
    feed = {"data": [{"id": "99_1", "created_time": "2026-10-15T08:00:00+0000", "message": "Hello", "permalink_url": "https://facebook.com/99/posts/1", "shares": {"count": 7},
                      "reactions": {"summary": {"total_count": 300}}, "comments": {"summary": {"total_count": 12}}}, {"id": "99_2", "created_time": "2026-10-14T08:00:00+0000"}]}
    got = Facebook(KEYS, opener([("/99/posts", feed), ("graph.facebook.com/v21.0/aajtak", page)])).fetch("aajtak")
    assert got.account.followers == 5_000_000 and got.account.external_id == "99"
    assert (got.posts[0].likes, got.posts[0].comments, got.posts[0].shares) == (300, 12, 7) and (got.posts[1].likes, got.posts[1].shares) == (None, None)


def test_facebook_page_without_post_permission_still_gives_followers_but_a_bad_key_is_reported():
    page = {"id": "99", "name": "P", "fan_count": 4_900_000}
    got = Facebook(KEYS, opener([("/99/posts", 403), ("graph.facebook.com/v21.0/p", page)])).fetch("p")
    assert got.account.followers == 4_900_000 and got.posts == []
    with pytest.raises(SocialError) as err:
        Facebook(KEYS, opener([("/99/posts", 401), ("graph.facebook.com/v21.0/p", page)])).fetch("p")
    assert err.value.code == "auth"


def test_x_account_and_posts():
    user = {"data": {"id": "42", "name": "StatAxis", "public_metrics": {"followers_count": 1000, "following_count": 10, "tweet_count": 77}}}
    tl = {"data": [{"id": "t1", "created_time": None, "created_at": "2026-10-15T06:00:00.000Z", "text": "hello", "public_metrics": {"like_count": 10, "reply_count": 2, "retweet_count": 3, "quote_count": 1, "impression_count": 900}},
                   {"id": "t2", "created_at": "2026-10-14T06:00:00.000Z", "text": "x", "public_metrics": {"like_count": 1}}]}
    got = X(KEYS, opener([("/users/42/tweets", tl), ("/users/by/username/StatAxis", user)])).fetch("StatAxis")
    assert (got.account.followers, got.account.following, got.account.posts_total) == (1000, 10, 77)
    assert (got.posts[0].likes, got.posts[0].comments, got.posts[0].shares, got.posts[0].views) == (10, 2, 4, 900)
    assert (got.posts[1].shares, got.posts[1].views, got.posts[1].comments) == (None, None, None)
    with pytest.raises(SocialError) as err:
        X(KEYS, opener([("/users/by/username/", {})])).fetch("nobody")
    assert err.value.code == "not_found"


# ------------------------------------------------------------------------------------------------------------------------------ accounts and reading
def test_accounts_are_validated_and_not_doubled():
    s, _ = make_session()
    a, b = channel(s, "a", "Alpha"), channel(s, "b", "Beta")
    acc = service.save_account(s, channel_id=a.id, platform="Twitter", handle="https://x.com/Alpha_News/", by="admin")
    assert (acc.platform, acc.handle) == ("x", "Alpha_News")
    with pytest.raises(ValueError):
        service.save_account(s, channel_id=a.id, platform="tiktok", handle="x", by="admin")
    with pytest.raises(ValueError):
        service.save_account(s, channel_id=a.id, platform="x", handle="not valid!", by="admin")
    with pytest.raises(LookupError):
        service.save_account(s, channel_id=999, platform="x", handle="ok", by="admin")
    other = service.save_account(s, channel_id=a.id, platform="instagram", handle="alpha", by="admin")
    with pytest.raises(ValueError):                                        # an account belongs to one channel only
        service.save_account(s, channel_id=b.id, platform="instagram", handle="ALPHA", by="admin", account_id=acc.id)
    again = service.save_account(s, channel_id=b.id, platform="x", handle="alpha_news", by="admin")
    assert again.id == acc.id and again.channel_id == b.id and s.query(SocialAccount).count() == 2 and other.id != acc.id


def fake_connectors(readings, ready=("instagram", "facebook", "x")):
    """readings: {(platform, handle): Fetched | SocialError}."""
    class C:
        def __init__(self, platform):
            self.platform = platform
            self.calls = []

        def ready(self):
            return self.platform in ready

        def fetch(self, handle, external_id=None):
            self.calls.append(handle)
            got = readings[(self.platform, handle)]
            if isinstance(got, Exception):
                raise got
            return got
    return {p: C(p) for p in ("instagram", "facebook", "x")}


def reading(followers, posts=()):
    from collector.social.base import AccountSnapshot, Fetched
    return Fetched(AccountSnapshot(followers=followers, following=5, posts_total=100, display_name="N", external_id="9"), list(posts))


def post(pid, when, likes=None, comments=None, shares=None, views=None, text="t"):
    from collector.social.base import PostSnapshot
    return PostSnapshot(post_id=pid, posted_at=when, text=text, permalink=f"https://p/{pid}", media_type="post", likes=likes, comments=comments, shares=shares, views=views)


def test_a_pass_reads_due_accounts_stores_posts_and_waits_for_the_interval():
    s, _ = make_session()
    a = channel(s, "a", "Alpha")
    acc = service.save_account(s, channel_id=a.id, platform="instagram", handle="alpha", by="t")
    conns = fake_connectors({("instagram", "alpha"): reading(1000, [post("p1", T0 - timedelta(days=1), likes=50, comments=5)])})
    out = service.collect_social(s, now=T0, connector_map=conns)
    assert out["read"] == 1 and out["new_posts"] == 1
    s.refresh(acc)
    assert acc.last_status == "ok" and acc.external_id == "9" and s.query(SocialStats).count() == 1 and s.query(SocialPostStats).count() == 1
    assert service.collect_social(s, now=T0 + timedelta(hours=2), connector_map=conns)["read"] == 0                    # read again only after the interval (6 hours)
    out = service.collect_social(s, now=T0 + timedelta(hours=7), connector_map=conns)
    assert out["read"] == 1 and out["new_posts"] == 0 and s.query(SocialPost).count() == 1 and s.query(SocialPostStats).count() == 2 and s.query(SocialStats).count() == 2
    assert service.collect_social(s, now=T0 + timedelta(hours=8), connector_map=conns, force=True)["read"] == 1


def test_a_platform_without_a_key_is_marked_not_an_error():
    s, _ = make_session()
    a = channel(s, "a", "Alpha")
    acc = service.save_account(s, channel_id=a.id, platform="x", handle="alpha", by="t")
    out = service.collect_social(s, now=T0, connector_map=fake_connectors({}, ready=()))
    assert out["no_key"] == 1 and out["read"] == 0 and out["failed"] == 0
    s.refresh(acc)
    assert acc.last_status == "not_configured" and acc.last_checked_at is None            # it is tried again as soon as a key is saved


def test_a_refused_key_or_rate_limit_stops_that_platform_only_and_x_is_limited():
    s, _ = make_session()
    a = channel(s, "a", "Alpha")
    for p, h in (("instagram", "i1"), ("instagram", "i2"), ("facebook", "f1"), ("x", "x1"), ("x", "x2"), ("x", "x3")):
        service.save_account(s, channel_id=a.id, platform=p, handle=h, by="t")
    conns = fake_connectors({("instagram", "i1"): SocialError("auth", "refused"), ("instagram", "i2"): reading(1), ("facebook", "f1"): reading(10),
                             ("x", "x1"): reading(1), ("x", "x2"): reading(2), ("x", "x3"): reading(3)})
    s.add(AppSetting(key="social.x_max_per_run", value="2"))
    s.commit()
    out = service.collect_social(s, now=T0, connector_map=conns)
    assert conns["instagram"].calls == ["i1"]                                              # the refused key is not tried on the second account
    assert out["read"] == 3 and out["failed"] == 1 and out["skipped_x_limit"] == 1 and conns["facebook"].calls == ["f1"]
    limited = service.collect_social(s, now=T0 + timedelta(hours=9), connector_map=fake_connectors({("instagram", "i1"): reading(1), ("instagram", "i2"): reading(1), ("facebook", "f1"): reading(1),
                                                                                                    ("x", "x1"): reading(1), ("x", "x2"): reading(1), ("x", "x3"): reading(1)}), limit=2)
    assert limited["read"] == 2 and limited["left_for_next_pass"] >= 1


# ------------------------------------------------------------------------------------------------------------------------------ sheets
def test_sheets_bring_in_accounts_figures_and_posts_and_say_which_rows_failed():
    s, _ = make_session()
    a = channel(s, "UCaaa", "Alpha", handle="@alphanews")
    csv_accounts = "channel,platform,handle\n@alphanews,instagram,@alpha_ig\nAlpha,facebook,https://facebook.com/alphafb\nNoSuch,x,nobody\nUCaaa,x,bad handle!\nUCaaa,x,alphax\n"
    out = service.import_accounts(s, csv_accounts, "admin")
    assert out["added"] == 3 and len(out["errors"]) == 2 and {e["row"] for e in out["errors"]} == {4, 5}
    assert out["errors"][0]["error"].startswith("channel 'NoSuch'")
    stats = "platform,handle,date,followers,following,posts\ninstagram,alpha_ig,2026-10-01,1000,10,50\ninstagram,alpha_ig,2026-10-15,1100,10,55\nx,alphax,2026-10-15,abc,,\nx,ghost,2026-10-15,5,,\ninstagram,alpha_ig,2030-01-01,5,,\n"
    out = service.import_stats(s, stats)
    assert out["stored"] == 2 and len(out["errors"]) == 3
    assert s.query(SocialStats).count() == 2 and {r.source for r in s.query(SocialStats)} == {"import"}
    again = service.import_stats(s, "platform;handle;date;followers\ninstagram;alpha_ig;15/10/2026;1111\n")                 # the same day is replaced, not doubled
    assert again["stored"] == 1 and s.query(SocialStats).count() == 2 and s.query(SocialStats).order_by(SocialStats.observed_at.desc()).first().followers == 1111
    posts = "platform,handle,post_id,posted_at,text,likes,comments,shares,views\ninstagram,alpha_ig,p1,2026-10-14,Hello world,100,10,,\ninstagram,alpha_ig,,2026-10-14,no id,1,1,1,1\nx,alphax,t9,2026-10-14T05:00:00Z,x,5,0,0,900\n"
    out = service.import_posts(s, posts, now=T0)
    assert out["posts"] == 2 and out["new_posts"] == 2 and len(out["errors"]) == 1
    p = s.query(SocialPost).filter_by(platform_post_id="p1").one()
    st = s.query(SocialPostStats).filter_by(post_id=p.id).one()
    assert (st.likes, st.comments, st.shares, st.views) == (100, 10, None, None)
    with pytest.raises(ValueError):
        service.import_stats(s, "x" * 2_100_000)


def test_a_figure_can_be_typed_in():
    s, _ = make_session()
    a = channel(s, "a", "Alpha")
    acc = service.save_account(s, channel_id=a.id, platform="facebook", handle="alpha", by="t")
    assert service.add_manual_stats(s, acc.id, day="2026-10-10", followers=5000, by="t")["date"] == "2026-10-10"
    with pytest.raises(ValueError):
        service.add_manual_stats(s, acc.id, day="2026-10-10", followers="")
    with pytest.raises(ValueError):
        service.add_manual_stats(s, acc.id, day="not a date", followers=1)
    assert s.query(SocialStats).one().source == "manual"


# ------------------------------------------------------------------------------------------------------------------------------ the rankings
def stats(s, acc, when, followers):
    s.add(SocialStats(account_id=acc.id, observed_at=when, followers=followers, source="api"))


def seeded():
    s, engine = make_session()
    a, b, c = channel(s, "a", "Alpha"), channel(s, "b", "Beta", language="English"), channel(s, "c", "Gamma", segment="business")
    ia = service.save_account(s, channel_id=a.id, platform="instagram", handle="alpha", by="t")
    ib = service.save_account(s, channel_id=b.id, platform="instagram", handle="beta", by="t")
    xa = service.save_account(s, channel_id=a.id, platform="x", handle="alpha", by="t")
    service.save_account(s, channel_id=c.id, platform="instagram", handle="gamma", by="t")           # no reading yet
    for days, ai, bi in ((40, 800, 4000), (31, 900, 4100), (10, 950, 4200), (0, 1000, 4400)):
        stats(s, ia, T0 - timedelta(days=days), ai)
        stats(s, ib, T0 - timedelta(days=days), bi)
    stats(s, xa, T0 - timedelta(hours=3), 700)
    s.commit()
    # Alpha's Instagram posts in the last 30 days: 6 posts, 3 with figures
    for i, (l, cm, sh) in enumerate([(100, 10, None), (50, 0, 0), (30, 5, 5), (None, None, None), (None, None, None), (None, None, None)]):
        p = SocialPost(account_id=ia.id, platform_post_id=f"a{i}", posted_at=T0 - timedelta(days=i + 1), text=f"post {i}", permalink=f"https://i/{i}", api_refreshed_at=T0)
        s.add(p)
        s.flush()
        if l is not None:
            s.add(SocialPostStats(post_id=p.id, observed_at=T0, likes=l, comments=cm, shares=sh))
    pb = SocialPost(account_id=ib.id, platform_post_id="b0", posted_at=T0 - timedelta(days=2), text="beta post", api_refreshed_at=T0)
    s.add(pb)
    s.flush()
    for i in range(5):
        q = SocialPost(account_id=ib.id, platform_post_id=f"bb{i}", posted_at=T0 - timedelta(days=3 + i), text="x", api_refreshed_at=T0)
        s.add(q)
        s.flush()
        s.add(SocialPostStats(post_id=q.id, observed_at=T0, likes=200, comments=20, shares=10, views=5000))
    s.add(SocialPostStats(post_id=pb.id, observed_at=T0, likes=400, comments=40, shares=0))
    s.add(ChannelStats(channel_id=a.id, observed_at=T0, subscribers=20000))
    s.commit()
    return s, engine, (a, b, c), (ia, ib, xa)


def test_growth_and_engagement_are_worked_out_honestly():
    s, _, (a, b, c), (ia, ib, xa) = seeded()
    rows = {(r["channel"], r["platform"]): r for r in social.account_rows(s, days=30, now=T0)}
    r = rows[("Alpha", "instagram")]
    assert r["followers"] == 1000 and r["growth_abs"] == 100 and r["growth_pct"] == 11.11 and r["growth_span_days"] == 31.0           # newest reading at or before 30 days ago was the 31-day-old one (900)... 
    assert r["posts_in_period"] == 6 and r["posts_measured"] == 3 and r["avg_interactions"] == 66.7 and r["low_sample"] is True      # (110 + 50 + 40) / 3
    assert r["engagement_rate"] == 6.667                                                                                              # 66.7 as a percentage of 1000 followers
    assert rows[("Beta", "instagram")]["low_sample"] is False and rows[("Beta", "instagram")]["posts_measured"] == 6
    x = rows[("Alpha", "x")]
    assert x["followers"] == 700 and x["growth_abs"] is None and x["growth_pct"] is None            # one reading: no growth, not zero
    g = rows[("Gamma", "instagram")]
    assert g["followers"] is None and g["engagement_rate"] is None and g["posts_in_period"] == 0


def test_the_ranking_is_per_platform_and_filters_work():
    s, *_ = seeded()
    out = social.overview(s, days=30, now=T0)
    ig = next(p for p in out["platforms"] if p["platform"] == "instagram")
    assert [r["channel"] for r in ig["rows"]] == ["Beta", "Alpha", "Gamma"] and [r["rank"] for r in ig["rows"]] == [1, 2, None]
    assert ig["followers_total"] == 5400 and ig["accounts"] == 3 and ig["with_data"] == 2
    assert out["coverage"] == {"channels": 3, "channels_with_an_account": 3, "accounts": 4, "accounts_with_data": 3}
    by_growth = social.overview(s, platform="instagram", days=30, sort="growth", now=T0)["platforms"][0]["rows"]
    assert by_growth[0]["channel"] == "Alpha" and by_growth[0]["growth_pct"] > by_growth[1]["growth_pct"]
    assert [r["channel"] for r in social.overview(s, platform="instagram", language="English", now=T0)["platforms"][0]["rows"]] == ["Beta"]
    assert [r["channel"] for r in social.overview(s, platform="instagram", segment="business", now=T0)["platforms"][0]["rows"]] == ["Gamma"]
    with pytest.raises(ValueError):
        social.overview(s, platform="tiktok")
    with pytest.raises(ValueError):
        social.overview(s, sort="loudest")


def test_presence_sets_youtube_beside_the_platforms():
    s, _, (a, b, c), _ = seeded()
    out = social.presence(s, days=30, now=T0)
    alpha = next(r for r in out["rows"] if r["channel"] == "Alpha")
    assert alpha["youtube"] == 20000 and alpha["instagram"] == 1000 and alpha["x"] == 700 and alpha["facebook"] is None and alpha["total"] == 21700 and alpha["platforms_covered"] == 3
    assert out["rows"][0]["channel"] == "Alpha" and next(r for r in out["rows"] if r["channel"] == "Gamma")["total"] is None


def test_one_channel_compare_and_best_posts():
    s, _, (a, b, c), (ia, ib, xa) = seeded()
    one = social.channel_social(s, a.id, days=90, now=T0)
    assert {r["platform"] for r in one["accounts"]} == {"instagram", "x"} and one["top_posts"][0]["interactions"] == 110 and one["posts_in_period"] == 6
    assert next(r for r in one["accounts"] if r["platform"] == "instagram")["series"][-1]["value"] == 1000 and one["youtube"]["subscribers"] == 20000
    assert one["by_hour"] and all(0 <= h["hour"] < 24 for h in one["by_hour"])
    with pytest.raises(LookupError):
        social.channel_social(s, 9999)
    cmp = social.compare(s, [a.id, b.id], platform="instagram", days=60, now=T0)
    assert [r["channel"] for r in cmp["rows"]] == ["Beta", "Alpha"] and len(cmp["rows"][0]["series"]) == 4
    with pytest.raises(ValueError):
        social.compare(s, [], platform="instagram")
    with pytest.raises(ValueError):
        social.compare(s, list(range(1, 14)), platform="instagram")
    top = social.top_posts(s, days=30, limit=3, now=T0)["posts"]
    assert top[0]["interactions"] == 440 and top[0]["channel"] == "Beta" and [p["rank"] for p in top] == [1, 2, 3]
    by_views = social.top_posts(s, days=30, sort="views", now=T0)["posts"]
    assert by_views[0]["views"] == 5000 and all(p["views"] is not None for p in by_views)


def test_coverage_counts_what_is_linked_and_what_fails():
    s, _, _, (ia, ib, xa) = seeded()
    ia.last_status, xa.last_status = "ok", "auth"
    s.commit()
    cov = social.coverage(s)
    assert cov["accounts"] == 4 and cov["by_platform"]["x"]["failing"] == 1 and cov["by_platform"]["instagram"]["with_data"] == 2 and cov["settings"]["meta_token"]["set"] is False


# ------------------------------------------------------------------------------------------------------------------------------ settings and routes
def user(s, email, plan="sx_enterprise", admin=False):
    u = User(email=email, password_hash=hash_password("a long enough pass"), plan=plan, requested_plan=plan, is_admin=admin, active=True, approval_status="approved", full_name=email, organization="Org")
    s.add(u)
    s.commit()
    return u


def caller(engine):
    app = social_application(lambda: Session(engine))

    def call(who, method, path, body=None, query=""):
        raw = json.dumps(body).encode() if body is not None else b""
        out = {}
        env = {"REQUEST_METHOD": method, "PATH_INFO": path, "QUERY_STRING": query, "wsgi.input": io.BytesIO(raw), "CONTENT_LENGTH": str(len(raw)),
               "HTTP_AUTHORIZATION": ("Bearer " + issue_token(AuthIdentity(who.id, who.email, who.plan, who.is_admin))) if who else ""}
        res = app(env, lambda status, headers: out.update(status=status, headers=dict(headers)))
        data = b"".join(res)
        ctype = out["headers"].get("Content-Type", "")
        return int(out["status"].split()[0]), (json.loads(data) if ctype.startswith("application/json") else data)
    return call


def test_keys_are_kept_masked_and_checked():
    s, _ = make_session()
    social_settings.save(s, "admin", {"meta_token": "EAAB" + "x" * 40, "meta_ig_user_id": "17841400000000000", "x_max_per_run": 10})
    st = social_settings.public_status(s)
    assert st["meta_token"]["hint"] == "…" + "x" * 4 and "EAAB" not in json.dumps(st) and st["meta_ig_user_id"]["hint"] == "17841400000000000" and st["x_max_per_run"] == 10 and st["x_bearer"]["set"] is False
    social_settings.save(s, "admin", {"meta_token": ""})                                               # an empty value leaves the saved key alone
    assert social_settings.credentials(s).meta_token.startswith("EAAB")
    social_settings.save(s, "admin", {"clear_meta_token": True})
    assert social_settings.credentials(s).meta_token == ""
    for bad in ({"meta_token": "short"}, {"meta_ig_user_id": "abc"}, {"x_max_per_run": "-1"}, {"hours_between": 0}):
        with pytest.raises(ValueError):
            social_settings.save(s, "admin", bad)


def test_the_routes_follow_the_privileges():
    s, engine, (a, b, c), _ = seeded()
    free, intel, admin = user(s, "free@x.com", "sx_free"), user(s, "intel@x.com", "sx_intelligence"), user(s, "root@x.com", "sx_enterprise", admin=True)
    call = caller(engine)
    assert call(None, "GET", "/api/v1/social/overview")[0] == 401
    assert call(free, "GET", "/api/v1/social/overview")[0] == 403
    status, body = call(intel, "GET", "/api/v1/social/overview", query="platform=instagram&days=30")
    assert status == 200 and body["platforms"][0]["platform"] == "instagram"
    assert call(intel, "GET", "/api/v1/social/overview", query="platform=tiktok")[0] == 400
    assert call(intel, "GET", "/api/v1/social/presence")[0] == 200 and call(intel, "GET", "/api/v1/social/posts", query="days=7")[0] == 200
    assert call(intel, "GET", "/api/v1/social/channel", query=f"channel_id={a.id}")[0] == 200 and call(intel, "GET", "/api/v1/social/channel", query="channel_id=99999")[0] == 404
    assert call(intel, "GET", "/api/v1/social/compare", query=f"channels={a.id},{b.id}&platform=instagram")[0] == 200 and call(intel, "GET", "/api/v1/social/compare", query="channels=a&platform=x")[0] == 400
    assert call(intel, "GET", "/api/v1/social/export.xlsx")[0] == 403                                       # the Excel file is the data-export privilege (SX Analyst and above)
    status, data = call(admin, "GET", "/api/v1/social/export.xlsx", query="platform=instagram")
    assert status == 200
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Accounts", "Top posts", "Notes"] and wb["Accounts"].max_row == 4 and wb["Accounts"]["A1"].value == "Channel"
    assert call(intel, "GET", "/api/v1/social/nothing")[0] == 404
    assert call(intel, "POST", "/api/v1/social/overview", {})[0] == 404


def test_only_an_administrator_can_set_it_up():
    s, engine, (a, b, c), _ = seeded()
    intel, admin = user(s, "intel@x.com", "sx_intelligence"), user(s, "root@x.com", "sx_enterprise", admin=True)
    call = caller(engine)
    for method, path, body in (("GET", "/api/v1/social/admin/status", None), ("POST", "/api/v1/social/admin/settings", {"x_max_per_run": 5}), ("POST", "/api/v1/social/admin/accounts/save", {"channel_id": a.id, "platform": "x", "handle": "q"}),
                               ("POST", "/api/v1/social/admin/run", {})):
        assert call(intel, method, path, body)[0] == 403, path
    status, st = call(admin, "GET", "/api/v1/social/admin/status")
    assert status == 200 and st["accounts"] == 4 and len(st["accounts_list"]) == 4 and st["settings"]["x_bearer"]["set"] is False
    status, saved = call(admin, "POST", "/api/v1/social/admin/settings", {"x_bearer": "AAAA" + "b" * 40, "x_max_per_run": 7})
    assert status == 200 and saved["settings"]["x_bearer"]["hint"] == "…" + "b" * 4 and saved["settings"]["x_max_per_run"] == 7 and "AAAA" not in json.dumps(saved)
    status, made = call(admin, "POST", "/api/v1/social/admin/accounts/save", {"channel_id": b.id, "platform": "facebook", "handle": "https://facebook.com/betafb"})
    assert status == 200 and made["account"]["handle"] == "betafb" and made["account"]["channel"] == "Beta"
    assert call(admin, "POST", "/api/v1/social/admin/accounts/save", {"channel_id": b.id, "platform": "tiktok", "handle": "x"})[0] == 400
    status, imp = call(admin, "POST", "/api/v1/social/admin/accounts/import", {"csv": "channel,platform,handle\nBeta,x,betax\n"})
    assert status == 200 and imp["added"] == 1
    status, st2 = call(admin, "POST", "/api/v1/social/admin/stats/add", {"account_id": made["account"]["id"], "date": "2026-10-10", "followers": 321})
    assert status == 200 and st2["stored"] is True
    status, run = call(admin, "POST", "/api/v1/social/admin/accounts/check", {"id": made["account"]["id"]})
    assert status == 200 and run["result"]["status"] == "not_configured"                                    # no Meta key saved: explained, not an error
    status, listing = call(admin, "GET", "/api/v1/social/admin/accounts", query="problems=1")
    assert status == 200 and all(a_["status"] not in (None, "ok") for a_ in listing["accounts"])
    status, gone = call(admin, "POST", "/api/v1/social/admin/accounts/delete", {"id": made["account"]["id"]})
    assert status == 200 and s.query(SocialAccount).filter_by(handle="betafb").count() == 0
    assert call(admin, "POST", "/api/v1/social/admin/accounts/delete", {"id": 99999})[0] == 404


# ------------------------------------------------------------------------------------------------------------------------------ keeping figures
def test_old_figures_and_stale_post_text_are_removed():
    s, _, (a, b, c), (ia, ib, xa) = seeded()
    s.add(SocialStats(account_id=ia.id, observed_at=T0 - timedelta(days=36 * 31), followers=1))
    old = SocialPost(account_id=ia.id, platform_post_id="old", posted_at=T0 - timedelta(days=36 * 31), text="old", api_refreshed_at=T0 - timedelta(days=36 * 31 + 5))
    s.add(old)
    s.flush()
    s.add(SocialPostStats(post_id=old.id, observed_at=T0 - timedelta(days=36 * 31), likes=1))
    stale = SocialPost(account_id=ia.id, platform_post_id="stale", posted_at=T0 - timedelta(days=60), text="stale text", permalink="https://x", api_refreshed_at=T0 - timedelta(days=45))
    s.add(stale)
    s.commit()
    out = run_social_retention(s, now=T0)
    assert out["stats_deleted"] == 1 and out["post_stats_deleted"] == 1 and out["posts_deleted"] == 1 and out["post_texts_cleared"] == 1
    assert s.query(SocialPost).filter_by(platform_post_id="old").count() == 0
    kept = s.query(SocialPost).filter_by(platform_post_id="stale").one()
    assert kept.text is None and kept.permalink is None
    assert s.query(SocialPost).filter_by(platform_post_id="a0").one().text == "post 0"                  # fresh posts keep their text


def test_the_demo_shows_invented_social_accounts_only():
    from api import demo_data
    s, _ = make_session()
    demo_data.seed(s, now=T0, days=2)
    out = social.overview(s, days=7, now=T0)
    assert out["coverage"]["accounts"] > 20 and out["coverage"]["accounts_with_data"] == out["coverage"]["accounts"]
    rows = [r for p in out["platforms"] for r in p["rows"]]
    assert all(r["handle"].startswith("sample_") and r["url"] is None and "invented" in r["display_name"] for r in rows)
    assert any(r["growth_abs"] and r["growth_abs"] > 0 for r in rows) and social.top_posts(s, days=7, now=T0)["posts"]


def test_the_social_part_of_a_periodic_report():
    from api import report_layouts
    s, _, _, _ = seeded()
    block = social.report_block(s, days=30, now=T0)
    assert [p["platform"] for p in block["platforms"]] == ["Instagram", "X"] and block["platforms"][0]["top"][0]["name"] == "Beta" and block["platforms"][0]["total"] == 5400
    empty, _ = make_session()
    assert social.report_block(empty) is None
    content = report_layouts.sample_content("weekly")
    for with_social in (None, block):
        content["social"] = with_social
        for build in (report_layouts.one_pager, report_layouts.booklet):
            assert build(content, sections=["data", "social"]).startswith(b"%PDF")
    assert "social" in report_layouts.ALL_SECTIONS and "social" not in report_layouts.ONE_PAGER_SECTIONS
