"""The Excel file of the social-media figures: one sheet of accounts (with growth and engagement) and one of posts. Machine-readable, so it is the data-export privilege, not the report privilege."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from api import social
from api.stamps import file_stamp, ist_cell, ist_stamp


def _sheet(workbook: Workbook, title: str, head: list[str], rows: list[list[Any]], first: bool = False) -> None:
    ws = workbook.active if first else workbook.create_sheet()
    ws.title = title
    ws.append(head)
    for row in rows:
        ws.append(row)
    for i, name in enumerate(head, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(48, len(name) + 4))
    ws.freeze_panes = "A2"


def export_social_xlsx(session: Session, *, platform: str | None, language: str | None, segment: str | None, days: int, now: datetime | None = None) -> tuple[bytes, str]:
    now = now or datetime.now(UTC)
    overview = social.overview(session, platform=platform, language=language, segment=segment, days=days, now=now)
    posts = social.top_posts(session, platform=platform, language=language, segment=segment, days=min(days, 90), limit=100, now=now)
    workbook = Workbook()
    accounts = [r for p in overview["platforms"] for r in p["rows"]]
    _sheet(workbook, "Accounts", ["Channel", "Language", "Category", "Platform", "Handle", "Followers", "Following", "Posts in total", "Read at (IST)", "Source", f"Growth over {days} days", f"Growth % over {days} days", "Days the growth covers",
                                  "Posts in period", "Posts with figures", "Average reactions per post", "Engagement rate % of followers", "Low sample", "Status", "Address"],
           [[r["channel"], r["language"], r["segment"], r["platform_label"], r["handle"], r["followers"], r["following"], r["posts_total"], ist_cell(r["observed_at"]), r["source"], r["growth_abs"], r["growth_pct"], r["growth_span_days"],
             r["posts_in_period"], r["posts_measured"], r["avg_interactions"], r["engagement_rate"], "yes" if r["low_sample"] else "no", r["status"], r["url"]] for r in accounts], first=True)
    _sheet(workbook, "Top posts", ["Rank", "Channel", "Platform", "Handle", "Posted (IST)", "Text (start)", "Likes", "Comments", "Shares", "Views", "Reactions", "Reactions % of followers", "Address"],
           [[p["rank"], p["channel"], p["platform_label"], p["handle"], ist_cell(p["posted_at"]), p["text"], p["likes"], p["comments"], p["shares"], p["views"], p["interactions"], p["rate"], p["permalink"]] for p in posts["posts"]])
    notes = workbook.create_sheet("Notes")
    for line in (["Downloaded at (IST)", ist_stamp(now)], ["Method", social.METHOD], ["Data", "Public figures read through each platform's official API, or supplied by the channel or an administrator. Empty means not given, not zero."]):
        notes.append(line)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 120
    buf = BytesIO()
    workbook.save(buf)
    return buf.getvalue(), f"stataxis-social-media-downloaded-{file_stamp(now)}.xlsx"
