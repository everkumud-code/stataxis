"""Role-aware Excel exports for StatAxis dashboard reports and live audience windows."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO

from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from sqlalchemy import Select, and_, func, select
from sqlalchemy.orm import Session

from api.stamps import file_stamp, ist_cell, ist_stamp
from api.access import UserRole, plan_video_access_policy, require_capability, video_access_policy
from api.plans import SXPlan, get_plan
from api.live_monitor import live_audience_window
from api.live_stats import live_window_stats
from collector.storage import Channel, Observation, Video
from metrics.persistence import IntelligenceSnapshotRecord


MAX_EXPORT_ROWS = 50_000


@dataclass(frozen=True)
class ObservationExportFilters:
    """Optional dashboard filters for an observation export."""

    start_at: datetime | None = None
    end_at: datetime | None = None
    language: str | None = None
    region: str | None = None
    channel_id: int | None = None
    video_id: int | None = None
    is_live: bool | None = None
    classification: str | None = None


def _apply_filters(stmt: Select[Any], filters: ObservationExportFilters) -> Select[Any]:
    if filters.start_at is not None:
        stmt = stmt.where(Observation.observed_at >= filters.start_at)
    if filters.end_at is not None:
        stmt = stmt.where(Observation.observed_at <= filters.end_at)
    if filters.language is not None:
        stmt = stmt.where(Channel.language == filters.language)
    if filters.region is not None:
        stmt = stmt.where(Channel.region == filters.region)
    if filters.channel_id is not None:
        stmt = stmt.where(Observation.channel_id == filters.channel_id)
    if filters.video_id is not None:
        stmt = stmt.where(Observation.video_id == filters.video_id)
    if filters.is_live is not None:
        stmt = stmt.where(Observation.is_live == filters.is_live)
    if filters.classification is not None:
        stmt = stmt.where(Observation.classification == filters.classification)
    return stmt


def _joined(*columns: Any) -> Select[Any]:
    return (select(*columns).select_from(Observation).join(Video, Video.id == Observation.video_id).join(Channel, Channel.id == Observation.channel_id)
            .where(Channel.active.is_(True))   # a channel that is switched off is not in any report or download
            .order_by(Observation.observed_at.asc(), Observation.id.asc()))


def filtered_observations(session: Session, filters: ObservationExportFilters) -> list[tuple[Observation, Video, Channel]]:
    """Return observations matching the dashboard export filters."""
    stmt: Select[tuple[Observation, Video, Channel]] = _apply_filters(_joined(Observation, Video, Channel), filters)
    rows = list(session.execute(stmt.limit(MAX_EXPORT_ROWS + 1)).all())
    if len(rows) > MAX_EXPORT_ROWS:
        raise ValueError("report has more than 50000 rows; narrow the date range or add filters")
    return rows


def _widths(rows: list[list[Any]], header: list[str]) -> list[int]:
    """Column widths from the header and the first few hundred rows (looking at every cell of a large sheet costs more than writing it)."""
    return [min(max([len(h)] + [len(str(r[i])) for r in rows[:300] if r[i] is not None]) + 2, 48) for i, h in enumerate(header)]


def _sheet(workbook: Workbook, title: str, header: list[str], rows: list[list[Any]]) -> None:
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    for i, width in enumerate(_widths(rows, header)):
        sheet.column_dimensions[get_column_letter(i + 1)].width = width
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"


def export_observations_xlsx(session: Session, filters: ObservationExportFilters) -> bytes:
    """Build an Excel workbook containing filtered data and latest intelligence. Plain columns and a streaming workbook: a day of readings is tens of thousands of rows."""
    columns = (Observation.observed_at, Video.id, Video.youtube_video_id, Video.title, Channel.name, Channel.language, Channel.region, Observation.classification, Observation.is_live,
               Observation.view_count, Observation.like_count, Observation.comment_count, Observation.concurrent_viewers, Observation.source)
    got = session.execute(_apply_filters(_joined(*columns), filters).limit(MAX_EXPORT_ROWS + 1)).all()
    if len(got) > MAX_EXPORT_ROWS:
        raise ValueError("report has more than 50000 rows; narrow the date range or add filters")
    workbook = Workbook(write_only=True)
    header = ["Observed At (IST)", "Video ID", "YouTube Video ID", "Title", "Channel", "Language", "Region", "Classification", "Live", "Views", "Likes", "Comments", "Concurrent Viewers", "Source"]
    data_rows = [[ist_cell(r[0]), *r[1:]] for r in got]
    _sheet(workbook, "StatAxis Data", header, data_rows)

    # one line for each video: its newest score (only the newest snapshot of each video is read)
    videos: dict[int, tuple[str, str]] = {}
    for r in got:
        videos.setdefault(r[1], (r[2], r[3]))
    latest: dict[int, IntelligenceSnapshotRecord] = {}
    ids = sorted(videos)
    for i in range(0, len(ids), 800):
        chunk = ids[i:i + 800]
        newest = select(IntelligenceSnapshotRecord.video_id.label("v"), func.max(IntelligenceSnapshotRecord.generated_at).label("m")).where(IntelligenceSnapshotRecord.video_id.in_(chunk)).group_by(IntelligenceSnapshotRecord.video_id).subquery()
        for record in session.execute(select(IntelligenceSnapshotRecord).join(newest, and_(IntelligenceSnapshotRecord.video_id == newest.c.v, IntelligenceSnapshotRecord.generated_at == newest.c.m))).scalars():
            latest.setdefault(record.video_id, record)
    lines = []
    for video_id in videos:
        record = latest.get(video_id)
        if record is None:
            continue
        view = json.loads(record.view_json)
        lines.append([video_id, videos[video_id][0], videos[video_id][1], record.score, record.confidence, record.available_signals, ist_cell(record.generated_at),
                      view.get("view", view.get("summary", "")), json.dumps(json.loads(record.contributions_json), sort_keys=True)])
    _sheet(workbook, "StatAxis Intelligence", ["Video ID", "YouTube Video ID", "Title", "STX Index", "Confidence", "Available Signals", "Generated At (IST)", "StatAxis View", "Signal Contributions"], lines)

    _add_notice_sheet(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_live_stats_xlsx(
    session: Session,
    *,
    start_at: datetime,
    end_at: datetime,
    language: str | None = None,
    segment: str | None = None,
    bucket_seconds: int = 60,
) -> bytes:
    """Excel export of average / peak concurrent viewers by Primary, Secondary and All feed."""
    payload = live_window_stats(session, start_at=start_at, end_at=end_at, language=language, segment=segment, bucket_seconds=bucket_seconds)
    workbook = Workbook()
    channels = workbook.active
    channels.title = "Channels"
    channels.append([
        "Rank", "Market", "Channel", "Language", "All Avg", "All Peak", "Primary Avg", "Primary Peak",
        "Secondary Avg", "Secondary Peak", "Share % (All Avg)", "Streams Seen", "Coverage %", "Peak At (IST)",
    ])
    for item in payload["channels"]:
        feeds = item["feeds"]
        channels.append([
            item["rank"], item["market_label"], item["name"], item["language"],
            feeds["all"]["average"], feeds["all"]["peak"], feeds["primary"]["average"], feeds["primary"]["peak"],
            feeds["secondary"]["average"], feeds["secondary"]["peak"], item.get("share_percent"),
            item["streams_seen"], item["coverage_percent"], ist_cell(item["peak_at"]),
        ])
    _format_sheet(channels)
    markets = workbook.create_sheet("Markets")
    markets.append(["Market", "Channels", "Average Concurrent", "Peak Concurrent", "Headline"])
    for item in payload["markets"]:
        markets.append([item["label"], item["channel_count"], item["average"], item["peak"], item["headline"]])
    _format_sheet(markets)
    notes = workbook.create_sheet("Method")
    notes.append(["Window start (IST)", ist_cell(payload["start_at"])])
    notes.append(["Window end (IST)", ist_cell(payload["end_at"])])
    notes.append(["Grid seconds", payload["bucket_seconds"]])
    notes.append(["Primary threshold (hours live)", payload["primary_min_live_hours"]])
    notes.append(["Method", payload["method"]])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_live_audience_xlsx(
    session: Session,
    *,
    start_at: datetime,
    end_at: datetime,
    language: str | None = None,
) -> bytes:
    """Export the exact observed live audience window, including language and channel detail."""
    payload = live_audience_window(session, start_at=start_at, end_at=end_at, language=language)
    workbook = Workbook()

    timeline = workbook.active
    timeline.title = "Second-by-Second"
    timeline.append(["Observed At (IST)", "Hindi", "English", "Regional", "Unknown", "Total Concurrent"])
    for point in payload["timeline"]:
        timeline.append([
            ist_cell(point["observed_at"]), point.get("Hindi", 0), point.get("English", 0),
            point.get("Regional", 0), point.get("Unknown", 0), point.get("total_concurrent", 0),
        ])
    _format_sheet(timeline)

    markets = workbook.create_sheet("Language Markets")
    markets.append(["Language Group", "Channel Count", "Observations", "Current Concurrent", "Peak Concurrent", "Average Concurrent", "Metric Note"])
    for group, item in payload["languages"].items():
        markets.append([
            group, item["channel_count"], item["observations"], item["current_concurrent"],
            item["peak_concurrent"], item["average_concurrent"], item.get("metric_note", ""),
        ])
    _format_sheet(markets)

    channels = workbook.create_sheet("Channels")
    channels.append(["Channel ID", "Channel", "Language Group", "Language", "Current Concurrent", "Peak Concurrent", "Latest Observed At (IST)"])
    for item in payload["channels"]:
        channels.append([
            item["channel_id"], item["name"], item["language_group"], item["language"],
            item["current_concurrent"], item["peak_concurrent"], ist_cell(item["observed_at"]),
        ])
    _format_sheet(channels)

    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Window Start (IST)", ist_cell(payload["start_at"])])
    summary.append(["Window End (IST)", ist_cell(payload["end_at"])])
    summary.append(["Peak Concurrent", payload["overall"]["peak_concurrent"]])
    summary.append(["Current Concurrent", payload["overall"]["current_concurrent"]])
    summary.append(["Observed Seconds", payload["overall"]["observed_seconds"]])
    summary.append(["Live Channels", payload["overall"]["channel_count"]])
    summary.append(["Interpolation", "OFF — observed data only"])
    summary.append(["Language Filter", language or "All"])
    summary.append(["Sampling Definition", payload["sample_resolution"]])
    _format_sheet(summary)

    _add_notice_sheet(workbook)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_for_role(session: Session, role: UserRole | SXPlan | str, filters: ObservationExportFilters) -> bytes:
    """Authorize and export the dashboard report for an authenticated role or plan."""
    try:
        policy = plan_video_access_policy(get_plan(role).code)
    except ValueError:
        policy = video_access_policy(role)
    require_capability(policy, "can_download_report")
    return export_observations_xlsx(session, filters)


def export_response(session: Session, role: UserRole | SXPlan | str, filters: ObservationExportFilters) -> tuple[int, dict[str, str], bytes]:
    """Return an HTTP-ready report response with server-side role enforcement."""
    try:
        payload = export_for_role(session, role, filters)
    except PermissionError as exc:
        return 403, {"Content-Type": "application/json"}, json.dumps({"error": str(exc)}).encode()
    except ValueError as exc:
        return 400, {"Content-Type": "application/json"}, json.dumps({"error": str(exc)}).encode()

    return 200, {
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "Content-Disposition": f"attachment; filename=stataxis-report-downloaded-{file_stamp()}.xlsx",
    }, payload


DERIVED_METRICS_NOTICE = (
    "STX Index, STAX9 and other scores in this workbook are metrics generated independently by StatAxis. "
    "They are not sourced from, provided by or endorsed by YouTube. Underlying statistics come from public "
    "YouTube API data."
)


def _add_notice_sheet(workbook: Workbook) -> None:
    notes = workbook.create_sheet("Notes")
    notes.column_dimensions["A"].width = 120   # set before the rows (a streaming workbook needs it so)
    notes.append(["Downloaded at", ist_stamp()])
    notes.append([])
    notes.append(["Data source", "Data via YouTube. Processed and evaluated by StatAxis."])
    notes.append(["Data notice"])
    notes.append([DERIVED_METRICS_NOTICE])


def _format_sheet(sheet) -> None:
    """Apply simple spreadsheet usability defaults."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 48)
        sheet.column_dimensions[column[0].column_letter].width = width
